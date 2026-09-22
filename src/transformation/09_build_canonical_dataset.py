from pathlib import Path
from openpyxl import load_workbook
from collections import defaultdict, Counter
import csv
import math
import re
import sys


# ============================================================
# PATHS
# ============================================================

ROOT = Path(__file__).resolve().parents[2]

RAW_DIR = ROOT / "row_data"
DISCOVERY_DIR = ROOT / "outputs" / "discovery"

CANONICAL_DIR = ROOT / "data" / "processed" / "canonical"
VALIDATION_DIR = ROOT / "outputs" / "validation"
DOC_ARCH = ROOT / "docs" / "architecture"
DOC_VALIDATION = ROOT / "docs" / "validation"

CANONICAL_DIR.mkdir(parents=True, exist_ok=True)
VALIDATION_DIR.mkdir(parents=True, exist_ok=True)
DOC_ARCH.mkdir(parents=True, exist_ok=True)
DOC_VALIDATION.mkdir(parents=True, exist_ok=True)

SOURCE_CATALOG = DISCOVERY_DIR / "11_four_year_comparability_candidates.csv"

FACT_CAPACITY = CANONICAL_DIR / "fact_capacity.csv"
FACT_ACTIVITY = CANONICAL_DIR / "fact_activity.csv"
FACT_WORKFORCE_SECTOR = CANONICAL_DIR / "fact_workforce_sector.csv"
FACT_WORKFORCE_NATIONALITY = CANONICAL_DIR / "fact_workforce_nationality.csv"

DIM_YEAR = CANONICAL_DIR / "dim_year.csv"
DIM_GEOGRAPHY = CANONICAL_DIR / "dim_geography.csv"
DIM_SECTOR = CANONICAL_DIR / "dim_sector.csv"
DIM_WORKFORCE = CANONICAL_DIR / "dim_workforce_type.csv"
DIM_NATIONALITY = CANONICAL_DIR / "dim_nationality.csv"

LINEAGE_FILE = VALIDATION_DIR / "source_lineage.csv"
PROFILE_FILE = VALIDATION_DIR / "canonical_profile.csv"
RECON_FILE = VALIDATION_DIR / "validation_reconciliation.csv"
WORKFORCE_RECON = VALIDATION_DIR / "workforce_cross_source_reconciliation.csv"
EXCLUDED_FILE = VALIDATION_DIR / "excluded_source_register.csv"

CONTRACT_MD = DOC_ARCH / "CANONICAL_DATA_CONTRACT.md"
SUMMARY_MD = DOC_VALIDATION / "PHASE_3_CANONICALIZATION_SUMMARY.md"

YEARS = [2021, 2022, 2023, 2024]

CORE_GROUPS = {
    "FY-005",
    "FY-006",
    "FY-008",
    "FY-010",
    "FY-018",
    "FY-020",
    "FY-023",
}


# ============================================================
# CANONICAL MAPPINGS
# ============================================================

SECTOR_MAP = {
    "وزارة الصحة": "Ministry of Health",
    "الجهات الحكومية الأخرى": "Other Governmental Sector",
    "الجهات الحكومية الاخرى": "Other Governmental Sector",
    "القطاع الخاص": "Private Sector",
}

WORKFORCE_MAP = {
    "أطباء بشريون": "Physicians",
    "أطباء أسنان": "Dentists",
    "طاقم التمريض": "Nurses",
    "قابلات": "Midwives",
    "صيادلة": "Pharmacists",
    "فئات طبية مساعدة": "Allied Health Personnel",
}

NATIONALITY_MAP = {
    "سعودي": "Saudi",
    "غير سعودي": "Non-Saudi",
}

ADMIN_REGION_MAP = {
    "الرياض": "Riyadh",
    "مكة المكرمة": "Makkah",
    "المدينة المنورة": "Madinah",
    "القصيم": "Qassim",
    "الشرقية": "Eastern Region",
    "عسير": "Asir",
    "تبوك": "Tabuk",
    "حائل": "Hail",
    "الحدود الشمالية": "Northern Borders",
    "جازان": "Jazan",
    "نجران": "Najran",
    "الباحة": "Al Baha",
    "الجوف": "Al Jouf",
}


# ============================================================
# HELPERS
# ============================================================

def clean(value):
    if value is None:
        return ""

    text = str(value).strip()
    text = text.replace("\n", " ")
    text = text.replace("\r", " ")

    return re.sub(r"\s+", " ", text).strip()


def norm_ar(value):
    text = clean(value)

    text = (
        text
        .replace("إ", "ا")
        .replace("أ", "ا")
        .replace("آ", "ا")
    )

    return re.sub(r"\s+", " ", text).strip()


def read_csv(path):
    if not path.exists():
        print(f"ERROR: Missing file: {path}")
        sys.exit(1)

    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path, rows, fields=None):

    if fields is None:
        fields = list(rows[0].keys()) if rows else []

    with path.open("w", encoding="utf-8-sig", newline="") as f:

        writer = csv.DictWriter(
            f,
            fieldnames=fields
        )

        writer.writeheader()

        if rows:
            writer.writerows(rows)


def to_number(value):

    if value is None or value == "":
        return None

    if isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return value

    text = clean(value).replace(",", "")

    # Formula = not a source numeric value
    if text.startswith("="):
        return None

    try:
        result = float(text)

        if result.is_integer():
            return int(result)

        return result

    except Exception:
        return None


def canonical_sector(value):

    value = clean(value)

    for source, canonical in SECTOR_MAP.items():
        if norm_ar(source) == norm_ar(value):
            return canonical

    return None


def canonical_workforce(value):

    value = clean(value)

    for source, canonical in WORKFORCE_MAP.items():
        if norm_ar(source) == norm_ar(value):
            return canonical

    return None


def canonical_nationality(value):

    value = clean(value)

    for source, canonical in NATIONALITY_MAP.items():
        if norm_ar(source) == norm_ar(value):
            return canonical

    return None


def canonical_region(value):

    value = clean(value)

    for source, canonical in ADMIN_REGION_MAP.items():
        if norm_ar(source) == norm_ar(value):
            return canonical

    return None


def find_year_cell(ws, year):

    target = f"{year}G"

    for row in ws.iter_rows(
        min_row=1,
        max_row=min(ws.max_row or 0, 20),
        values_only=False
    ):
        for cell in row:
            if clean(cell.value) == target:
                return cell.row, cell.column

    raise ValueError(
        f"Year header {target} not found in sheet {ws.title}"
    )


def get_catalog_source(catalog, group_id, year):

    candidates = [
        row for row in catalog
        if row["comparison_group_id"] == group_id
        and int(row["source_year"]) == year
    ]

    if not candidates:
        raise ValueError(
            f"Source not found for {group_id} / {year}"
        )

    return candidates[0]


# ============================================================
# LOAD CATALOG
# ============================================================

catalog = read_csv(SOURCE_CATALOG)

catalog = [
    row for row in catalog
    if row["comparison_group_id"] in CORE_GROUPS
]


# ============================================================
# WORKBOOK CACHE
# ============================================================

wb_formula_cache = {}
wb_value_cache = {}


def get_ws(source):
    """
    Return two versions of the same worksheet:

    1. Formula worksheet:
       data_only=False
       Preserves formulas and source structure.

    2. Value worksheet:
       data_only=True
       Returns Excel cached calculated values when available.

    This is required because some official MOH workbook cells contain
    formulas or external workbook references rather than literal numbers.
    """

    workbook_name = source["workbook_name"]
    sheet_name = source["sheet_name"]

    path = RAW_DIR / workbook_name

    if not path.exists():
        raise FileNotFoundError(path)

    if workbook_name not in wb_formula_cache:

        wb_formula_cache[workbook_name] = load_workbook(
            path,
            read_only=True,
            data_only=False
        )

    if workbook_name not in wb_value_cache:

        wb_value_cache[workbook_name] = load_workbook(
            path,
            read_only=True,
            data_only=True
        )

    wb_formula = wb_formula_cache[workbook_name]
    wb_value = wb_value_cache[workbook_name]

    if sheet_name not in wb_formula.sheetnames:
        raise KeyError(
            f"{sheet_name} not found in {workbook_name}"
        )

    if sheet_name not in wb_value.sheetnames:
        raise KeyError(
            f"{sheet_name} not found in cached-value workbook "
            f"{workbook_name}"
        )

    return (
        wb_formula[sheet_name],
        wb_value[sheet_name]
    )


def source_number(ws_formula, ws_value, row, column):
    """
    Extract a numeric source value safely.

    Priority:
    1. Literal numeric value from formula-preserving worksheet.
    2. Cached Excel result from data_only=True worksheet.

    Missing values remain None.
    They are never converted to zero.
    """

    formula_value = ws_formula.cell(row, column).value

    direct_number = to_number(formula_value)

    if direct_number is not None:
        return direct_number

    cached_value = ws_value.cell(row, column).value

    cached_number = to_number(cached_value)

    return cached_number


# ============================================================
# OUTPUT COLLECTIONS
# ============================================================

fact_capacity = []
fact_activity = []
fact_workforce_sector = []
fact_workforce_nationality = []

lineage = []
issues = []


def add_issue(
    severity,
    group,
    year,
    issue,
    detail
):
    issues.append({
        "severity": severity,
        "source_group": group,
        "year": year,
        "check": issue,
        "status": "FAIL" if severity == "ERROR" else "REVIEW",
        "details": detail
    })


# ============================================================
# FY-005
# HOSPITALS / BEDS / BEDS PER 10,000
#
# Grain detail:
# Year × National × Sector × Capacity Measure
# ============================================================

def parse_fy005(year):

    source = get_catalog_source(
        catalog,
        "FY-005",
        year
    )

    ws, ws_values = get_ws(source)

    year_row, year_col = find_year_cell(
        ws,
        year
    )

    data_start = year_row + 2

    extracted = 0

    for r in range(
        data_start,
        min(ws.max_row, data_start + 15) + 1
    ):

        source_sector = clean(
            ws.cell(r, 1).value
        )

        sector = canonical_sector(
            source_sector
        )

        if sector:

            hospitals = source_number(
                ws,
                ws_values,
                r,
                year_col
            )

            beds = source_number(
                ws,
                ws_values,
                r,
                year_col + 2
            )

            if hospitals is None or beds is None:
                add_issue(
                    "ERROR",
                    "FY-005",
                    year,
                    "NUMERIC_EXTRACTION",
                    f"Missing numeric Hospitals/Beds for {sector}"
                )
                continue

            for measure, value, unit in [
                ("Hospitals", hospitals, "Count"),
                ("Beds", beds, "Count"),
            ]:

                fact_capacity.append({
                    "Year": year,
                    "GeographyType": "National",
                    "Geography": "Saudi Arabia",
                    "Sector": sector,
                    "CapacityMeasure": measure,
                    "Value": value,
                    "Unit": unit,
                    "SourceGroup": "FY-005",
                    "SourceWorkbook": source["workbook_name"],
                    "SourceSheet": source["sheet_name"]
                })

            extracted += 1

        # National rate
        if (
            "معدل" in source_sector
            and "عشرة" in source_sector
        ):

            rate = source_number(
                ws,
                ws_values,
                r,
                year_col + 2
            )

            if rate is not None:

                fact_capacity.append({
                    "Year": year,
                    "GeographyType": "National",
                    "Geography": "Saudi Arabia",
                    "Sector": "All Sectors",
                    "CapacityMeasure": "Beds per 10,000 population",
                    "Value": rate,
                    "Unit": "Beds per 10,000 population",
                    "SourceGroup": "FY-005",
                    "SourceWorkbook": source["workbook_name"],
                    "SourceSheet": source["sheet_name"]
                })

    if extracted != 3:
        add_issue(
            "ERROR",
            "FY-005",
            year,
            "EXPECTED_SECTOR_COUNT",
            f"Expected 3 sectors; extracted {extracted}"
        )


# ============================================================
# FY-006
# WORKFORCE BY HEALTH SECTOR
#
# Grain:
# Year × Sector × Workforce Type
# ============================================================

def parse_fy006(year):

    source = get_catalog_source(
        catalog,
        "FY-006",
        year
    )

    ws, ws_values = get_ws(source)

    year_row, year_col = find_year_cell(
        ws,
        year
    )

    data_start = year_row + 2

    current_sector = None
    extracted = 0

    for r in range(
        data_start,
        min(ws.max_row, data_start + 50) + 1
    ):

        raw_sector = clean(
            ws.cell(r, 1).value
        )

        mapped_sector = canonical_sector(
            raw_sector
        )

        if mapped_sector:
            current_sector = mapped_sector

        # Stop before grand total sector
        if norm_ar(raw_sector) in {
            norm_ar("المجموع"),
            norm_ar("المجموع الكلي")
        }:
            break

        raw_category = clean(
            ws.cell(r, 2).value
        )

        workforce = canonical_workforce(
            raw_category
        )

        if (
            current_sector
            and workforce
        ):

            count_value = source_number(
                ws,
                ws_values,
                r,
                year_col
            )

            saudi_percent = source_number(
                ws,
                ws_values,
                r,
                year_col + 1
            )

            if count_value is None:
                add_issue(
                    "ERROR",
                    "FY-006",
                    year,
                    "NUMERIC_EXTRACTION",
                    f"No workforce count for "
                    f"{current_sector} / {workforce}"
                )
                continue

            fact_workforce_sector.append({
                "Year": year,
                "GeographyType": "National",
                "Geography": "Saudi Arabia",
                "Sector": current_sector,
                "WorkforceType": workforce,
                "WorkforceCount": count_value,
                "SaudiPercent": saudi_percent,
                "Unit": "Persons",
                "SourceGroup": "FY-006",
                "SourceWorkbook": source["workbook_name"],
                "SourceSheet": source["sheet_name"]
            })

            extracted += 1

    expected = 3 * 6

    if extracted != expected:
        add_issue(
            "ERROR",
            "FY-006",
            year,
            "EXPECTED_GRAIN_ROWS",
            f"Expected {expected}; extracted {extracted}"
        )


# ============================================================
# FY-008 / FY-010
# WORKFORCE BY NATIONALITY
#
# Grain:
# Year × Scope × Workforce Type × Nationality
# ============================================================

def parse_nationality_group(
    group_id,
    year,
    scope
):

    source = get_catalog_source(
        catalog,
        group_id,
        year
    )

    ws, ws_values = get_ws(source)

    year_row, year_col = find_year_cell(
        ws,
        year
    )

    data_start = year_row + 1

    current_workforce = None
    extracted = 0

    for r in range(
        data_start,
        min(ws.max_row, data_start + 45) + 1
    ):

        raw_category = clean(
            ws.cell(r, 1).value
        )

        mapped_workforce = canonical_workforce(
            raw_category
        )

        if mapped_workforce:
            current_workforce = mapped_workforce

        raw_nationality = clean(
            ws.cell(r, 2).value
        )

        nationality = canonical_nationality(
            raw_nationality
        )

        if (
            current_workforce
            and nationality
        ):

            value = source_number(
                ws,
                ws_values,
                r,
                year_col
            )

            if value is None:
                add_issue(
                    "ERROR",
                    group_id,
                    year,
                    "NUMERIC_EXTRACTION",
                    f"No value for "
                    f"{scope} / {current_workforce} / "
                    f"{nationality}"
                )
                continue

            fact_workforce_nationality.append({
                "Year": year,
                "GeographyType": "National",
                "Geography": "Saudi Arabia",
                "Scope": scope,
                "WorkforceType": current_workforce,
                "Nationality": nationality,
                "WorkforceCount": value,
                "Unit": "Persons",
                "SourceGroup": group_id,
                "SourceWorkbook": source["workbook_name"],
                "SourceSheet": source["sheet_name"]
            })

            extracted += 1

    expected = 6 * 2

    if extracted != expected:
        add_issue(
            "ERROR",
            group_id,
            year,
            "EXPECTED_GRAIN_ROWS",
            f"Expected {expected}; extracted {extracted}"
        )


# ============================================================
# FY-018
# RED CRESCENT BY ADMINISTRATIVE REGION
#
# Capacity:
# Year × Administrative Region × Capacity Measure
#
# Activity:
# Year × Administrative Region × Activity Measure
# ============================================================

def parse_fy018(year):

    source = get_catalog_source(
        catalog,
        "FY-018",
        year
    )

    ws, ws_values = get_ws(source)

    extracted_regions = set()

    for r in range(1, ws.max_row + 1):

        raw_region = clean(
            ws.cell(r, 1).value
        )

        region = canonical_region(
            raw_region
        )

        if not region:
            continue

        cases = source_number(
            ws,
            ws_values,
            r,
            2
        )

        centers = source_number(
            ws,
            ws_values,
            r,
            3
        )

        ambulances = source_number(
            ws,
            ws_values,
            r,
            5
        )

        if None in (cases, centers, ambulances):
            add_issue(
                "ERROR",
                "FY-018",
                year,
                "NUMERIC_EXTRACTION",
                f"Missing source values for {region}"
            )
            continue

        extracted_regions.add(region)

        for measure, value in [
            ("First Aid Centers", centers),
            ("Ambulances", ambulances),
        ]:

            fact_capacity.append({
                "Year": year,
                "GeographyType": "Administrative Region",
                "Geography": region,
                "Sector": "Saudi Red Crescent Authority",
                "CapacityMeasure": measure,
                "Value": value,
                "Unit": "Count",
                "SourceGroup": "FY-018",
                "SourceWorkbook": source["workbook_name"],
                "SourceSheet": source["sheet_name"]
            })

        fact_activity.append({
            "Year": year,
            "GeographyType": "Administrative Region",
            "Geography": region,
            "Sector": "Saudi Red Crescent Authority",
            "ActivityMeasure":
                "Cases offered first aid / transported to hospitals",
            "Value": cases,
            "Unit": "Cases",
            "SourceGroup": "FY-018",
            "SourceWorkbook": source["workbook_name"],
            "SourceSheet": source["sheet_name"]
        })

    if len(extracted_regions) != 13:
        add_issue(
            "ERROR",
            "FY-018",
            year,
            "EXPECTED_REGION_COUNT",
            f"Expected 13 administrative regions; "
            f"found {len(extracted_regions)}"
        )


# ============================================================
# FY-020
# ENCOUNTERS BY SECTOR + NATIONAL ENCOUNTERS PER PERSON
# ============================================================

def parse_fy020(year):

    source = get_catalog_source(
        catalog,
        "FY-020",
        year
    )

    ws, ws_values = get_ws(source)

    year_row, year_col = find_year_cell(
        ws,
        year
    )

    data_start = year_row + 1

    sectors_found = 0
    rate_found = False

    for r in range(
        data_start,
        min(ws.max_row, data_start + 12) + 1
    ):

        label = clean(
            ws.cell(r, 1).value
        )

        sector = canonical_sector(
            label
        )

        if sector:

            value = source_number(
                ws,
                ws_values,
                r,
                year_col
            )

            if value is None:
                add_issue(
                    "ERROR",
                    "FY-020",
                    year,
                    "NUMERIC_EXTRACTION",
                    f"Missing encounters for {sector}"
                )
                continue

            fact_activity.append({
                "Year": year,
                "GeographyType": "National",
                "Geography": "Saudi Arabia",
                "Sector": sector,
                "ActivityMeasure": "Encounters",
                "Value": value,
                "Unit": "Encounters",
                "SourceGroup": "FY-020",
                "SourceWorkbook": source["workbook_name"],
                "SourceSheet": source["sheet_name"]
            })

            sectors_found += 1

        if (
            "متوسط" in label
            and "زيارات" in label
        ):

            rate = source_number(
                ws,
                ws_values,
                r,
                year_col
            )

            if rate is not None:

                fact_activity.append({
                    "Year": year,
                    "GeographyType": "National",
                    "Geography": "Saudi Arabia",
                    "Sector": "All Sectors",
                    "ActivityMeasure":
                        "Encounters per person per year",
                    "Value": rate,
                    "Unit": "Encounters per person",
                    "SourceGroup": "FY-020",
                    "SourceWorkbook": source["workbook_name"],
                    "SourceSheet": source["sheet_name"]
                })

                rate_found = True

    if sectors_found != 3:
        add_issue(
            "ERROR",
            "FY-020",
            year,
            "EXPECTED_SECTOR_COUNT",
            f"Expected 3 sectors; extracted {sectors_found}"
        )

    if not rate_found:
        add_issue(
            "ERROR",
            "FY-020",
            year,
            "RATE_EXTRACTION",
            "Encounters per person rate was not found"
        )


# ============================================================
# FY-023
# INPATIENTS + ADMISSIONS PER 100 PERSONS
# ============================================================

def parse_fy023(year):

    source = get_catalog_source(
        catalog,
        "FY-023",
        year
    )

    ws, ws_values = get_ws(source)

    year_row, year_col = find_year_cell(
        ws,
        year
    )

    data_start = year_row + 1

    sectors_found = 0
    rate_found = False

    for r in range(
        data_start,
        min(ws.max_row, data_start + 12) + 1
    ):

        label = clean(
            ws.cell(r, 1).value
        )

        sector = canonical_sector(
            label
        )

        if sector:

            value = source_number(
                ws,
                ws_values,
                r,
                year_col
            )

            if value is None:
                add_issue(
                    "ERROR",
                    "FY-023",
                    year,
                    "NUMERIC_EXTRACTION",
                    f"Missing inpatient value for {sector}"
                )
                continue

            fact_activity.append({
                "Year": year,
                "GeographyType": "National",
                "Geography": "Saudi Arabia",
                "Sector": sector,
                "ActivityMeasure":
                    "Inpatients / Admissions",
                "Value": value,
                "Unit": "Admissions",
                "SourceGroup": "FY-023",
                "SourceWorkbook": source["workbook_name"],
                "SourceSheet": source["sheet_name"]
            })

            sectors_found += 1

        if (
            "متوسط" in label
            and "100" in label
        ):

            rate = source_number(
                ws,
                ws_values,
                r,
                year_col
            )

            if rate is not None:

                fact_activity.append({
                    "Year": year,
                    "GeographyType": "National",
                    "Geography": "Saudi Arabia",
                    "Sector": "All Sectors",
                    "ActivityMeasure":
                        "Admissions per 100 persons",
                    "Value": rate,
                    "Unit": "Admissions per 100 persons",
                    "SourceGroup": "FY-023",
                    "SourceWorkbook": source["workbook_name"],
                    "SourceSheet": source["sheet_name"]
                })

                rate_found = True

    if sectors_found != 3:
        add_issue(
            "ERROR",
            "FY-023",
            year,
            "EXPECTED_SECTOR_COUNT",
            f"Expected 3 sectors; extracted {sectors_found}"
        )

    if not rate_found:
        add_issue(
            "ERROR",
            "FY-023",
            year,
            "RATE_EXTRACTION",
            "Admissions rate was not found"
        )


# ============================================================
# EXECUTE EXTRACTION
# ============================================================

print("")
print("============================================================")
print("PHASE 3 — CANONICAL EXTRACTION")
print("============================================================")

for year in YEARS:

    print(f"Processing year {year}...")

    parse_fy005(year)
    parse_fy006(year)

    parse_nationality_group(
        "FY-008",
        year,
        "MOH Total"
    )

    parse_nationality_group(
        "FY-010",
        year,
        "MOH Primary Health Care Centers"
    )

    parse_fy018(year)
    parse_fy020(year)
    parse_fy023(year)



# ============================================================
# HARD CONTRACT VALIDATION — WORKFORCE NATIONALITY
#
# Expected grain:
# 4 years
# × 2 scopes
# × 6 workforce types
# × 2 nationalities
# = 96 rows
# ============================================================

expected_workforce_nationality_rows = (
    len(YEARS)
    * 2
    * len(WORKFORCE_MAP)
    * len(NATIONALITY_MAP)
)

actual_workforce_nationality_rows = len(
    fact_workforce_nationality
)

if (
    actual_workforce_nationality_rows
    != expected_workforce_nationality_rows
):
    raise RuntimeError(
        "WORKFORCE NATIONALITY CONTRACT FAILURE: "
        f"Expected {expected_workforce_nationality_rows} rows "
        f"but extracted {actual_workforce_nationality_rows}."
    )


expected_matrix = {
    (
        year,
        scope,
        workforce,
        nationality
    )
    for year in YEARS
    for scope in [
        "MOH Total",
        "MOH Primary Health Care Centers"
    ]
    for workforce in WORKFORCE_MAP.values()
    for nationality in NATIONALITY_MAP.values()
}


actual_matrix = {
    (
        int(row["Year"]),
        row["Scope"],
        row["WorkforceType"],
        row["Nationality"]
    )
    for row in fact_workforce_nationality
}


missing_matrix_rows = (
    expected_matrix - actual_matrix
)

unexpected_matrix_rows = (
    actual_matrix - expected_matrix
)


if missing_matrix_rows:
    raise RuntimeError(
        "WORKFORCE NATIONALITY CONTRACT FAILURE — "
        f"Missing grain rows: "
        f"{sorted(missing_matrix_rows)}"
    )


if unexpected_matrix_rows:
    raise RuntimeError(
        "WORKFORCE NATIONALITY CONTRACT FAILURE — "
        f"Unexpected grain rows: "
        f"{sorted(unexpected_matrix_rows)}"
    )


# Specific regression test for the source issue discovered
# in FY-008 / 2021.

regression_rows = [
    row
    for row in fact_workforce_nationality
    if (
        int(row["Year"]) == 2021
        and row["Scope"] == "MOH Total"
        and row["WorkforceType"] == "Pharmacists"
        and row["Nationality"] == "Non-Saudi"
    )
]


if len(regression_rows) != 1:
    raise RuntimeError(
        "REGRESSION FAILURE: Expected exactly one "
        "2021 MOH Total / Pharmacists / Non-Saudi row."
    )


regression_value = to_number(
    regression_rows[0]["WorkforceCount"]
)


if regression_value != 131:
    raise RuntimeError(
        "REGRESSION FAILURE: "
        "2021 MOH Total / Pharmacists / Non-Saudi "
        f"expected 131 but extracted {regression_value}."
    )


print(
    "PASS: Workforce nationality matrix = "
    f"{actual_workforce_nationality_rows} / "
    f"{expected_workforce_nationality_rows}"
)

print(
    "PASS: FY-008 2021 Non-Saudi Pharmacists = 131"
)

# ============================================================
# CLOSE SOURCE WORKBOOKS
# ============================================================

for wb in wb_formula_cache.values():
    try:
        wb.close()
    except Exception:
        pass

for wb in wb_value_cache.values():
    try:
        wb.close()
    except Exception:
        pass


# ============================================================
# SORT FACT TABLES
# ============================================================

fact_capacity.sort(
    key=lambda x: (
        x["Year"],
        x["GeographyType"],
        x["Geography"],
        x["Sector"],
        x["CapacityMeasure"]
    )
)

fact_activity.sort(
    key=lambda x: (
        x["Year"],
        x["GeographyType"],
        x["Geography"],
        x["Sector"],
        x["ActivityMeasure"]
    )
)

fact_workforce_sector.sort(
    key=lambda x: (
        x["Year"],
        x["Sector"],
        x["WorkforceType"]
    )
)

fact_workforce_nationality.sort(
    key=lambda x: (
        x["Year"],
        x["Scope"],
        x["WorkforceType"],
        x["Nationality"]
    )
)


# ============================================================
# WRITE FACTS
# ============================================================

write_csv(
    FACT_CAPACITY,
    fact_capacity
)

write_csv(
    FACT_ACTIVITY,
    fact_activity
)

write_csv(
    FACT_WORKFORCE_SECTOR,
    fact_workforce_sector
)

write_csv(
    FACT_WORKFORCE_NATIONALITY,
    fact_workforce_nationality
)


# ============================================================
# DIMENSIONS
# ============================================================

dim_year = [
    {
        "YearKey": year,
        "Year": year
    }
    for year in YEARS
]


geographies = {
    (
        row["GeographyType"],
        row["Geography"]
    )
    for row in (
        fact_capacity
        + fact_activity
        + fact_workforce_sector
        + fact_workforce_nationality
    )
}

dim_geography = []

for i, (geo_type, geography) in enumerate(
    sorted(geographies),
    start=1
):
    dim_geography.append({
        "GeographyKey": i,
        "GeographyType": geo_type,
        "Geography": geography
    })


sectors = sorted({
    row["Sector"]
    for row in (
        fact_capacity
        + fact_activity
        + fact_workforce_sector
    )
})

dim_sector = [
    {
        "SectorKey": i,
        "Sector": value
    }
    for i, value in enumerate(
        sectors,
        start=1
    )
]


workforce_types = sorted({
    row["WorkforceType"]
    for row in (
        fact_workforce_sector
        + fact_workforce_nationality
    )
})

dim_workforce = [
    {
        "WorkforceTypeKey": i,
        "WorkforceType": value
    }
    for i, value in enumerate(
        workforce_types,
        start=1
    )
]


nationalities = sorted({
    row["Nationality"]
    for row in fact_workforce_nationality
})

dim_nationality = [
    {
        "NationalityKey": i,
        "Nationality": value
    }
    for i, value in enumerate(
        nationalities,
        start=1
    )
]


write_csv(DIM_YEAR, dim_year)
write_csv(DIM_GEOGRAPHY, dim_geography)
write_csv(DIM_SECTOR, dim_sector)
write_csv(DIM_WORKFORCE, dim_workforce)
write_csv(DIM_NATIONALITY, dim_nationality)


# ============================================================
# LINEAGE
# ============================================================

lineage_seen = set()

for rows in [
    fact_capacity,
    fact_activity,
    fact_workforce_sector,
    fact_workforce_nationality
]:

    for row in rows:

        key = (
            row["SourceGroup"],
            row["Year"],
            row["SourceWorkbook"],
            row["SourceSheet"]
        )

        if key in lineage_seen:
            continue

        lineage_seen.add(key)

        lineage.append({
            "SourceGroup": row["SourceGroup"],
            "SourceYear": row["Year"],
            "SourceWorkbook": row["SourceWorkbook"],
            "SourceSheet": row["SourceSheet"],
            "SourceFolder": "row_data",
            "SourceModified": "NO"
        })


write_csv(
    LINEAGE_FILE,
    sorted(
        lineage,
        key=lambda x: (
            x["SourceGroup"],
            x["SourceYear"]
        )
    )
)


# ============================================================
# VALIDATION — DUPLICATES / NULLS / YEAR COVERAGE
# ============================================================

validation_rows = []


def validate_fact(
    name,
    rows,
    key_fields,
    value_fields
):

    keys = Counter(
        tuple(row[field] for field in key_fields)
        for row in rows
    )

    duplicates = sum(
        1
        for count in keys.values()
        if count > 1
    )

    null_values = 0

    for row in rows:
        for field in value_fields:
            if row.get(field) in (None, ""):
                null_values += 1

    years = sorted({
        int(row["Year"])
        for row in rows
    })

    validation_rows.append({
        "Check": f"{name} row count",
        "Status": "PASS",
        "Actual": len(rows),
        "Expected": "> 0",
        "Variance": "",
        "Details": ""
    })

    validation_rows.append({
        "Check": f"{name} duplicate grain",
        "Status": "PASS" if duplicates == 0 else "FAIL",
        "Actual": duplicates,
        "Expected": 0,
        "Variance": duplicates,
        "Details": "Duplicate natural keys"
    })

    validation_rows.append({
        "Check": f"{name} null measure values",
        "Status": "PASS" if null_values == 0 else "FAIL",
        "Actual": null_values,
        "Expected": 0,
        "Variance": null_values,
        "Details": ""
    })

    validation_rows.append({
        "Check": f"{name} year coverage",
        "Status":
            "PASS"
            if years == YEARS
            else "FAIL",
        "Actual": " | ".join(map(str, years)),
        "Expected": "2021 | 2022 | 2023 | 2024",
        "Variance": "",
        "Details": ""
    })


validate_fact(
    "fact_capacity",
    fact_capacity,
    [
        "Year",
        "GeographyType",
        "Geography",
        "Sector",
        "CapacityMeasure"
    ],
    ["Value"]
)

validate_fact(
    "fact_activity",
    fact_activity,
    [
        "Year",
        "GeographyType",
        "Geography",
        "Sector",
        "ActivityMeasure"
    ],
    ["Value"]
)

validate_fact(
    "fact_workforce_sector",
    fact_workforce_sector,
    [
        "Year",
        "Geography",
        "Sector",
        "WorkforceType"
    ],
    [
        "WorkforceCount",
        "SaudiPercent"
    ]
)

validate_fact(
    "fact_workforce_nationality",
    fact_workforce_nationality,
    [
        "Year",
        "Geography",
        "Scope",
        "WorkforceType",
        "Nationality"
    ],
    ["WorkforceCount"]
)


# Extraction issues
for issue in issues:
    validation_rows.append({
        "Check":
            f"{issue['source_group']} "
            f"{issue['year']} "
            f"{issue['check']}",
        "Status": issue["status"],
        "Actual": "",
        "Expected": "",
        "Variance": "",
        "Details": issue["details"]
    })


# ============================================================
# CROSS-SOURCE RECONCILIATION
#
# FY-006 MOH count / Saudi%
# versus
# FY-008 Saudi + Non-Saudi
# ============================================================

sector_index = {}

for row in fact_workforce_sector:

    if row["Sector"] == "Ministry of Health":

        sector_index[
            (
                int(row["Year"]),
                row["WorkforceType"]
            )
        ] = row


nationality_index = defaultdict(dict)

for row in fact_workforce_nationality:

    if row["Scope"] != "MOH Total":
        continue

    nationality_index[
        (
            int(row["Year"]),
            row["WorkforceType"]
        )
    ][row["Nationality"]] = row["WorkforceCount"]


workforce_recon_rows = []

for year in YEARS:

    for workforce in sorted(WORKFORCE_MAP.values()):

        sector_row = sector_index.get(
            (year, workforce)
        )

        nat = nationality_index.get(
            (year, workforce),
            {}
        )

        saudi = to_number(
            nat.get("Saudi")
        )

        non_saudi = to_number(
            nat.get("Non-Saudi")
        )

        if (
            sector_row is None
            or saudi is None
            or non_saudi is None
        ):
            status = "FAIL"

            workforce_recon_rows.append({
                "Year": year,
                "WorkforceType": workforce,
                "FY006_MOH_Count":
                    sector_row["WorkforceCount"]
                    if sector_row else "",
                "FY008_Saudi": saudi,
                "FY008_NonSaudi": non_saudi,
                "FY008_ReconstructedTotal": "",
                "CountVariance": "",
                "FY006_SaudiPercent":
                    sector_row["SaudiPercent"]
                    if sector_row else "",
                "FY008_ReconstructedSaudiPercent": "",
                "SaudiPercentVariance": "",
                "Status": status
            })

            continue

        reconstructed_total = (
            saudi + non_saudi
        )

        source_total = to_number(
            sector_row["WorkforceCount"]
        )

        source_pct = to_number(
            sector_row["SaudiPercent"]
        )

        reconstructed_pct = (
            saudi / reconstructed_total * 100
            if reconstructed_total
            else None
        )

        count_variance = (
            source_total
            - reconstructed_total
        )

        pct_variance = (
            source_pct
            - reconstructed_pct
            if source_pct is not None
            and reconstructed_pct is not None
            else None
        )

        count_pass = (
            abs(count_variance) < 0.000001
        )

        # Source Saudi% is usually rounded to 1 decimal
        pct_pass = (
            pct_variance is not None
            and abs(pct_variance) <= 0.11
        )

        status = (
            "PASS"
            if count_pass and pct_pass
            else "FAIL"
        )

        workforce_recon_rows.append({
            "Year": year,
            "WorkforceType": workforce,
            "FY006_MOH_Count": source_total,
            "FY008_Saudi": saudi,
            "FY008_NonSaudi": non_saudi,
            "FY008_ReconstructedTotal":
                reconstructed_total,
            "CountVariance": count_variance,
            "FY006_SaudiPercent": source_pct,
            "FY008_ReconstructedSaudiPercent":
                round(reconstructed_pct, 4),
            "SaudiPercentVariance":
                round(pct_variance, 4)
                if pct_variance is not None
                else "",
            "Status": status
        })


write_csv(
    WORKFORCE_RECON,
    workforce_recon_rows
)


recon_failures = sum(
    1
    for row in workforce_recon_rows
    if row["Status"] != "PASS"
)

validation_rows.append({
    "Check":
        "FY-006 vs FY-008 MOH Workforce Reconciliation",
    "Status":
        "PASS" if recon_failures == 0 else "FAIL",
    "Actual":
        f"{len(workforce_recon_rows) - recon_failures} PASS",
    "Expected":
        f"{len(workforce_recon_rows)} PASS",
    "Variance": recon_failures,
    "Details":
        "FY006 Ministry of Health workforce count and Saudi% "
        "reconciled independently against FY008 Saudi + Non-Saudi."
})


write_csv(
    RECON_FILE,
    validation_rows
)


# ============================================================
# EXCLUDED SOURCE REGISTER
# ============================================================

excluded = [
    {
        "SourceGroup": "FY-029",
        "Domain": "Blood Bank Activity",
        "Decision": "EXCLUDE_FROM_CORE_TREND",
        "Reason":
            "Measure definitions change materially across 2021–2024. "
            "2021 includes separate transfusion and infectious-disease "
            "investigations; 2022–2023 use total investigations and "
            "collected/transfused units; 2024 introduces blood donors "
            "and changes the measure set.",
        "FutureUse":
            "May be analyzed separately after measure-level comparability "
            "contract is approved."
    }
]

write_csv(
    EXCLUDED_FILE,
    excluded
)


# ============================================================
# CANONICAL PROFILE
# ============================================================

profile_rows = [
    {
        "Dataset": "fact_capacity",
        "Rows": len(fact_capacity),
        "Years": len(set(r["Year"] for r in fact_capacity)),
        "SourceGroups":
            " | ".join(
                sorted(set(r["SourceGroup"] for r in fact_capacity))
            )
    },
    {
        "Dataset": "fact_activity",
        "Rows": len(fact_activity),
        "Years": len(set(r["Year"] for r in fact_activity)),
        "SourceGroups":
            " | ".join(
                sorted(set(r["SourceGroup"] for r in fact_activity))
            )
    },
    {
        "Dataset": "fact_workforce_sector",
        "Rows": len(fact_workforce_sector),
        "Years": len(set(r["Year"] for r in fact_workforce_sector)),
        "SourceGroups":
            " | ".join(
                sorted(set(r["SourceGroup"] for r in fact_workforce_sector))
            )
    },
    {
        "Dataset": "fact_workforce_nationality",
        "Rows": len(fact_workforce_nationality),
        "Years": len(set(r["Year"] for r in fact_workforce_nationality)),
        "SourceGroups":
            " | ".join(
                sorted(
                    set(
                        r["SourceGroup"]
                        for r in fact_workforce_nationality
                    )
                )
            )
    },
]

write_csv(
    PROFILE_FILE,
    profile_rows
)


# ============================================================
# DATA CONTRACT
# ============================================================

with CONTRACT_MD.open(
    "w",
    encoding="utf-8"
) as f:

    f.write(
        "# Canonical Data Contract\n\n"
    )

    f.write(
        "## Project\n\n"
        "**Saudi Healthcare Capacity & Performance Analytics 2021–2024**\n\n"
    )

    f.write(
        "## Source Rule\n\n"
        "For the canonical annual dataset, each reporting year is taken "
        "from the Statistical Yearbook published for that same year. "
        "Historical rolling-year columns contained in later yearbooks are "
        "not loaded as duplicate observations.\n\n"
    )

    f.write(
        "Example:\n\n"
        "- 2021 → 2021 Yearbook\n"
        "- 2022 → 2022 Yearbook\n"
        "- 2023 → 2023 Yearbook\n"
        "- 2024 → 2024 Yearbook\n\n"
    )

    f.write(
        "This prevents duplicated historical observations from rolling "
        "five-year tables.\n\n"
    )

    f.write(
        "## Fact_Capacity Grain\n\n"
        "`One row = Year × Geography Type × Geography × Sector × Capacity Measure`\n\n"
    )

    f.write(
        "Current measure families:\n\n"
        "- Hospitals\n"
        "- Beds\n"
        "- Beds per 10,000 population\n"
        "- First Aid Centers\n"
        "- Ambulances\n\n"
    )

    f.write(
        "## Fact_Activity Grain\n\n"
        "`One row = Year × Geography Type × Geography × Sector × Activity Measure`\n\n"
    )

    f.write(
        "Current measure families:\n\n"
        "- Encounters\n"
        "- Encounters per person per year\n"
        "- Inpatients / Admissions\n"
        "- Admissions per 100 persons\n"
        "- Saudi Red Crescent cases offered first aid / transported\n\n"
    )

    f.write(
        "## Fact_Workforce_Sector Grain\n\n"
        "`One row = Year × National Geography × Health Sector × Workforce Type`\n\n"
    )

    f.write(
        "Measures:\n\n"
        "- WorkforceCount\n"
        "- SaudiPercent\n\n"
    )

    f.write(
        "## Fact_Workforce_Nationality Grain\n\n"
        "`One row = Year × National Geography × Scope × Workforce Type × Nationality`\n\n"
    )

    f.write(
        "Scopes:\n\n"
        "- MOH Total\n"
        "- MOH Primary Health Care Centers\n\n"
    )

    f.write(
        "Nationalities:\n\n"
        "- Saudi\n"
        "- Non-Saudi\n\n"
    )

    f.write(
        "## Workforce Types\n\n"
        "- Physicians\n"
        "- Dentists\n"
        "- Nurses\n"
        "- Midwives\n"
        "- Pharmacists\n"
        "- Allied Health Personnel\n\n"
    )

    f.write(
        "Aggregate categories such as `Physicians & Dentists`, "
        "`Total Nurses and Midwives`, and grand totals are intentionally "
        "excluded from canonical detail facts to avoid double counting.\n\n"
    )

    f.write(
        "## Geography Contract\n\n"
        "Current approved geography concepts:\n\n"
        "- National → Saudi Arabia\n"
        "- Administrative Region → 13 Saudi administrative regions\n\n"
    )

    f.write(
        "`Health Region` and `Health Cluster` are not currently loaded "
        "into the Core canonical model. They remain separate concepts for "
        "future extensions.\n\n"
    )

    f.write(
        "## Total Behavior\n\n"
        "Source grand-total rows are not loaded as detail rows where they "
        "would duplicate lower-grain observations. Power BI / SQL totals "
        "should aggregate the canonical detail rows unless the KPI contract "
        "specifically defines an official published rate.\n\n"
    )

    f.write(
        "Published rates such as `Beds per 10,000 population`, "
        "`Encounters per person`, and `Admissions per 100 persons` are "
        "kept as separate official measures and must not be summed across "
        "categories.\n\n"
    )

    f.write(
        "## FY-029 Decision\n\n"
        "Blood Bank Activity is excluded from the Core multi-year trend "
        "because the available measure definitions change materially "
        "between years. It may be analyzed separately under a dedicated "
        "measure-level contract.\n\n"
    )

    f.write(
        "## Null / Blank / Zero Rule\n\n"
        "- NULL / missing source value = unknown or unavailable\n"
        "- Blank label = structural / inherited header where applicable\n"
        "- Zero = genuine reported zero only\n"
        "- No missing numeric value is automatically converted to zero\n\n"
    )

    f.write(
        "## Validation Rule\n\n"
        "Canonical outputs require:\n\n"
        "- zero duplicate natural keys\n"
        "- complete 2021–2024 coverage\n"
        "- no missing primary measure values\n"
        "- source lineage retained\n"
        "- independent reconciliation where another official source table "
        "supports the same metric\n"
    )


# ============================================================
# SUMMARY
# ============================================================

failed_checks = [
    row
    for row in validation_rows
    if row["Status"] == "FAIL"
]

review_checks = [
    row
    for row in validation_rows
    if row["Status"] == "REVIEW"
]

with SUMMARY_MD.open(
    "w",
    encoding="utf-8"
) as f:

    f.write(
        "# Phase 3 — Canonicalization Summary\n\n"
    )

    f.write(
        "## Outputs\n\n"
    )

    for row in profile_rows:
        f.write(
            f"- `{row['Dataset']}`: "
            f"**{row['Rows']} rows**, "
            f"{row['Years']} years, "
            f"source groups: {row['SourceGroups']}\n"
        )

    f.write(
        "\n## Validation\n\n"
    )

    f.write(
        f"- Validation checks: **{len(validation_rows)}**\n"
    )

    f.write(
        f"- Failed checks: **{len(failed_checks)}**\n"
    )

    f.write(
        f"- Review checks: **{len(review_checks)}**\n"
    )

    f.write(
        f"- Workforce cross-source reconciliations: "
        f"**{len(workforce_recon_rows)}**\n"
    )

    f.write(
        f"- Workforce reconciliation failures: "
        f"**{recon_failures}**\n\n"
    )

    f.write(
        "## Source Protection\n\n"
        "**Raw Excel workbooks were not modified.**\n\n"
    )

    f.write(
        "## Next Gate\n\n"
        "Do not load the canonical facts into SQL until "
        "`validation_reconciliation.csv` and "
        "`workforce_cross_source_reconciliation.csv` are reviewed.\n"
    )


# ============================================================
# FINAL OUTPUT
# ============================================================

print("")
print("============================================================")
print("PHASE 3 COMPLETE")
print("============================================================")

print(f"fact_capacity rows:              {len(fact_capacity)}")
print(f"fact_activity rows:              {len(fact_activity)}")
print(f"fact_workforce_sector rows:      {len(fact_workforce_sector)}")
print(
    f"fact_workforce_nationality rows: "
    f"{len(fact_workforce_nationality)}"
)

print("")
print(f"Validation failures:             {len(failed_checks)}")
print(f"Validation review items:         {len(review_checks)}")
print(
    f"Workforce reconciliation fails: "
    f"{recon_failures}"
)

print("")
print("Raw source files modified:       NO")
print("FY-029 loaded into Core:         NO")

print("")
print("PRIMARY FILES TO UPLOAD:")
print(PROFILE_FILE)
print(RECON_FILE)
print(WORKFORCE_RECON)
print(FACT_CAPACITY)
print(FACT_ACTIVITY)
print(FACT_WORKFORCE_SECTOR)
print(FACT_WORKFORCE_NATIONALITY)
print(CONTRACT_MD)
print(SUMMARY_MD)


