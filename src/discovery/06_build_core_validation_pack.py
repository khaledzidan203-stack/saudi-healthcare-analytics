from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import csv
import re
import sys


ROOT = Path(__file__).resolve().parents[2]

RAW_DIR = ROOT / "row_data"
DISCOVERY_DIR = ROOT / "outputs" / "discovery"
OUTPUT_DIR = ROOT / "outputs" / "validation"
DOCS_DIR = ROOT / "docs" / "architecture"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
DOCS_DIR.mkdir(parents=True, exist_ok=True)

FOUR_YEAR_FILE = DISCOVERY_DIR / "11_four_year_comparability_candidates.csv"

PACK_FILE = OUTPUT_DIR / "CORE_SOURCE_VALIDATION_PACK.xlsx"
APPROVAL_FILE = OUTPUT_DIR / "FINAL_SOURCE_APPROVAL.csv"
MODEL_FILE = DOCS_DIR / "DATA_MODEL_CONTRACT_V1.md"


# ============================================================
# SELECTED CORE GROUPS
# ============================================================

SELECTED_GROUPS = [
    "FY-005",   # Hospitals & Beds / Capacity
    "FY-006",   # Health Manpower
    "FY-008",   # MOH Manpower / Nationality
    "FY-010",   # PHC Manpower
    "FY-018",   # Red Crescent Centers / Ambulances
    "FY-020",   # Visits / Population
    "FY-023",   # Inpatients
    "FY-029",   # Blood Bank Activities
]

PREVIEW_ROWS = 35
MAX_COLUMNS = 60


# ============================================================
# HELPERS
# ============================================================

def read_csv(path):
    if not path.exists():
        print(f"ERROR: Missing file: {path}")
        sys.exit(1)

    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def clean(value):
    if value is None:
        return ""

    value = str(value).strip()
    value = value.replace("\n", " ").replace("\r", " ")

    return re.sub(r"\s+", " ", value)


def safe_sheet_name(name):
    name = re.sub(r'[:\\/*?\[\]]', '-', name)
    return name[:31]


def extract_year(filename):
    m = re.search(r"(2021|2022|2023|2024)", filename)
    return int(m.group(1)) if m else None


# ============================================================
# LOAD SOURCE CATALOG
# ============================================================

rows = read_csv(FOUR_YEAR_FILE)

selected = [
    r for r in rows
    if r["comparison_group_id"] in SELECTED_GROUPS
]

if not selected:
    print("ERROR: Selected groups were not found.")
    sys.exit(1)


by_group = {}

for row in selected:
    by_group.setdefault(
        row["comparison_group_id"],
        []
    ).append(row)


# ============================================================
# CREATE VALIDATION WORKBOOK
# ============================================================

out_wb = Workbook()

default_ws = out_wb.active
out_wb.remove(default_ws)

index_ws = out_wb.create_sheet("INDEX")

headers = [
    "Group ID",
    "Topic",
    "Representative Title",
    "Years Found",
    "Workbook / Sheet Sources",
    "Business Meaning Confirmed",
    "Grain Confirmed",
    "Geography Confirmed",
    "Units Confirmed",
    "Totals Reviewed",
    "KEEP?",
    "Notes"
]

for col, value in enumerate(headers, start=1):
    cell = index_ws.cell(row=1, column=col, value=value)
    cell.font = Font(bold=True)

index_ws.freeze_panes = "A2"


source_wb_cache = {}

approval_rows = []

index_row = 2


for group_id in SELECTED_GROUPS:

    group_rows = by_group.get(group_id, [])

    if not group_rows:
        continue

    group_rows.sort(
        key=lambda x: int(x["source_year"])
    )

    topic = group_rows[0].get("topic", "")
    title = group_rows[0].get(
        "likely_table_title",
        group_rows[0].get("sheet_name", "")
    )

    ws_name = safe_sheet_name(
        f"{group_id}_{topic[:15]}"
    )

    ws = out_wb.create_sheet(ws_name)

    # --------------------------------------------------------
    # Header section
    # --------------------------------------------------------

    ws["A1"] = "Comparison Group"
    ws["B1"] = group_id

    ws["A2"] = "Topic"
    ws["B2"] = topic

    ws["A3"] = "Candidate Title"
    ws["B3"] = title

    ws["A4"] = "Validation Status"
    ws["B4"] = "MANUAL REVIEW REQUIRED"

    for cell in ["A1", "A2", "A3", "A4"]:
        ws[cell].font = Font(bold=True)

    current_col = 1

    source_descriptions = []
    years_found = []


    # ========================================================
    # YEAR-BY-YEAR SIDE-BY-SIDE SOURCE VIEW
    # ========================================================

    for source_row in group_rows:

        year = int(source_row["source_year"])
        workbook_name = source_row["workbook_name"]
        sheet_name = source_row["sheet_name"]

        years_found.append(str(year))

        source_descriptions.append(
            f"{year}: {workbook_name} -> {sheet_name}"
        )

        source_path = RAW_DIR / workbook_name

        if workbook_name not in source_wb_cache:

            source_wb_cache[workbook_name] = load_workbook(
                source_path,
                read_only=True,
                data_only=False
            )

        src_wb = source_wb_cache[workbook_name]

        if sheet_name not in src_wb.sheetnames:
            continue

        src_ws = src_wb[sheet_name]

        max_rows = min(
            src_ws.max_row or 0,
            PREVIEW_ROWS
        )

        max_cols = min(
            src_ws.max_column or 0,
            MAX_COLUMNS
        )

        # Year title
        year_cell = ws.cell(
            row=6,
            column=current_col,
            value=f"YEAR {year}"
        )

        year_cell.font = Font(
            bold=True,
            size=12
        )

        year_cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7"
        )

        ws.cell(
            row=7,
            column=current_col,
            value=workbook_name
        )

        ws.cell(
            row=8,
            column=current_col,
            value=sheet_name
        )

        ws.cell(
            row=9,
            column=current_col,
            value=(
                f"Reported size: "
                f"{src_ws.max_row} rows × "
                f"{src_ws.max_column} columns"
            )
        )

        # Raw preview
        output_row = 11

        for values in src_ws.iter_rows(
            min_row=1,
            max_row=max_rows,
            min_col=1,
            max_col=max_cols,
            values_only=True
        ):

            for offset, value in enumerate(values):

                ws.cell(
                    row=output_row,
                    column=current_col + offset,
                    value=value
                )

            output_row += 1


        # Label separator
        current_col += max_cols + 2


    # ========================================================
    # FORMAT SHEET
    # ========================================================

    ws.freeze_panes = "A11"

    for col in range(
        1,
        min(ws.max_column, 120) + 1
    ):
        ws.column_dimensions[
            get_column_letter(col)
        ].width = 15


    # ========================================================
    # INDEX ENTRY
    # ========================================================

    index_ws.cell(index_row, 1, group_id)
    index_ws.cell(index_row, 2, topic)
    index_ws.cell(index_row, 3, title)
    index_ws.cell(
        index_row,
        4,
        " | ".join(years_found)
    )
    index_ws.cell(
        index_row,
        5,
        "\n".join(source_descriptions)
    )

    index_ws.cell(index_row, 6, "PENDING")
    index_ws.cell(index_row, 7, "PENDING")
    index_ws.cell(index_row, 8, "PENDING")
    index_ws.cell(index_row, 9, "PENDING")
    index_ws.cell(index_row, 10, "PENDING")
    index_ws.cell(index_row, 11, "PENDING")
    index_ws.cell(index_row, 12, "")

    index_ws.cell(
        index_row,
        5
    ).alignment = Alignment(
        wrap_text=True,
        vertical="top"
    )


    approval_rows.append({
        "comparison_group_id": group_id,
        "topic": topic,
        "representative_title": title,
        "years_found": " | ".join(years_found),

        "business_definition_confirmed": "PENDING",
        "grain_confirmed": "PENDING",
        "geography_confirmed": "PENDING",
        "units_confirmed": "PENDING",
        "totals_subtotals_reviewed": "PENDING",
        "footnotes_reviewed": "PENDING",

        "approved_grain": "",
        "approved_geography": "",
        "approved_measure_family": "",
        "approved_fact_table": "",

        "final_source_decision": "PENDING",
        "decision_reason": "",
        "review_notes": ""
    })

    index_row += 1


# ============================================================
# INDEX FORMATTING
# ============================================================

for col in range(
    1,
    index_ws.max_column + 1
):
    index_ws.column_dimensions[
        get_column_letter(col)
    ].width = 23

index_ws.column_dimensions["C"].width = 45
index_ws.column_dimensions["E"].width = 70
index_ws.column_dimensions["L"].width = 50


# ============================================================
# SAVE PACK
# ============================================================

out_wb.save(PACK_FILE)

for wb in source_wb_cache.values():
    try:
        wb.close()
    except Exception:
        pass


# ============================================================
# FINAL SOURCE APPROVAL CSV
# ============================================================

fields = [
    "comparison_group_id",
    "topic",
    "representative_title",
    "years_found",

    "business_definition_confirmed",
    "grain_confirmed",
    "geography_confirmed",
    "units_confirmed",
    "totals_subtotals_reviewed",
    "footnotes_reviewed",

    "approved_grain",
    "approved_geography",
    "approved_measure_family",
    "approved_fact_table",

    "final_source_decision",
    "decision_reason",
    "review_notes"
]

with APPROVAL_FILE.open(
    "w",
    encoding="utf-8-sig",
    newline=""
) as f:

    writer = csv.DictWriter(
        f,
        fieldnames=fields
    )

    writer.writeheader()
    writer.writerows(approval_rows)


# ============================================================
# DATA MODEL CONTRACT V1
# ============================================================

with MODEL_FILE.open(
    "w",
    encoding="utf-8"
) as f:

    f.write(
        "# Saudi Healthcare Capacity & Performance Analytics 2021–2024\n"
    )

    f.write(
        "## Data Model Contract V1 — Draft for Manual Approval\n\n"
    )

    f.write(
        "> Status: **DRAFT — source validation pending**\n\n"
    )

    f.write("## Business Problem\n\n")

    f.write(
        "Understand how healthcare capacity, workforce, and service "
        "activity evolved across Saudi Arabia from 2021 to 2024, "
        "and identify geographic differences that warrant deeper "
        "capacity and performance review.\n\n"
    )

    f.write("## Portfolio Scope\n\n")

    f.write(
        "The project intentionally uses a focused subset of the official "
        "Saudi Ministry of Health Statistical Yearbooks rather than all "
        "available worksheets.\n\n"
    )

    f.write("Core domains:\n\n")
    f.write("1. Population context\n")
    f.write("2. Healthcare capacity\n")
    f.write("3. Healthcare workforce\n")
    f.write("4. Healthcare activity\n")
    f.write("5. Geographic comparison\n")
    f.write("6. Multi-year trend analysis\n\n")

    f.write("## Selected Source Groups for Validation\n\n")

    for row in approval_rows:
        f.write(
            f"- `{row['comparison_group_id']}` — "
            f"{row['topic']} — "
            f"{row['representative_title']}\n"
        )

    f.write("\n## Geography Contract\n\n")

    f.write(
        "The following concepts must remain distinct unless an official "
        "mapping proves equivalence:\n\n"
    )

    f.write("- National\n")
    f.write("- Health Region\n")
    f.write("- Administrative Region\n")
    f.write("- Health Cluster\n\n")

    f.write(
        "A Health Cluster must not automatically be treated as a Health "
        "Region merely because it contains the same city or regional name.\n\n"
    )

    f.write("## Preliminary Star Schema\n\n")

    f.write("```text\n")
    f.write("                  Dim_Year\n")
    f.write("                     |\n")
    f.write("                  Dim_Geography\n")
    f.write("                     |\n")
    f.write("   -----------------------------------------\n")
    f.write("   |                 |                     |\n")
    f.write("Fact_Capacity   Fact_Workforce       Fact_Activity\n")
    f.write("   |                 |                     |\n")
    f.write("Dim_Facility    Dim_WorkforceType    Dim_ActivityType\n")
    f.write("\n")
    f.write("Optional / separate:\n")
    f.write("Fact_PopulationIndicators\n")
    f.write("```\n\n")

    f.write(
        "This model is not approved until source grain validation is "
        "completed.\n\n"
    )

    f.write("## Grain Contract Template\n\n")

    f.write(
        "Every approved fact table must document:\n\n"
    )

    f.write("```text\n")
    f.write("One row = Year × Geography × Business Measure Dimension\n")
    f.write("```\n\n")

    f.write(
        "Actual grain may differ by domain and must follow the official "
        "source table structure.\n\n"
    )

    f.write("## KPI Families — Candidate Only\n\n")

    f.write("### Capacity\n")
    f.write("- Hospitals\n")
    f.write("- Beds\n")
    f.write("- PHC / healthcare facilities where supported\n")
    f.write("- Ambulances / emergency infrastructure where supported\n\n")

    f.write("### Workforce\n")
    f.write("- Physicians\n")
    f.write("- Nurses\n")
    f.write("- Pharmacists\n")
    f.write("- Allied healthcare workforce where supported\n\n")

    f.write("### Activity\n")
    f.write("- Visits\n")
    f.write("- Inpatients / admissions\n")
    f.write("- Selected healthcare activities\n\n")

    f.write("### Derived Rates\n")
    f.write(
        "- Population-normalized capacity/workforce ratios only when "
        "numerator and denominator geography and year are compatible\n"
    )
    f.write(
        "- Resource-to-activity ratios only when fact grains are "
        "analytically compatible\n\n"
    )

    f.write("## KPI Contract Rule\n\n")

    f.write(
        "Before implementation every KPI must define:\n\n"
    )

    f.write(
        "Business Definition → Numerator → Denominator → Grain → "
        "Population → Filters → Time Logic → Unit → Total Behavior → "
        "Validation Baseline\n\n"
    )

    f.write("## Validation Gate\n\n")

    f.write(
        "Cleaning and canonicalization may begin only after each selected "
        "source group has been reviewed for:\n\n"
    )

    f.write("- business meaning\n")
    f.write("- grain\n")
    f.write("- geography\n")
    f.write("- units\n")
    f.write("- totals/subtotals\n")
    f.write("- footnotes/source notes\n")
    f.write("- final KEEP / EXCLUDE decision\n\n")

    f.write("## Current Status\n\n")
    f.write("**Manual source validation required.**\n")


# ============================================================
# FINAL
# ============================================================

print("")
print("============================================================")
print("PHASE 2B COMPLETE")
print("============================================================")
print(f"Selected groups:           {len(approval_rows)}")
print("Source files modified:     NO")
print("Cleaning performed:        NO")

print("")
print("UPLOAD THESE 3 FILES:")
print(PACK_FILE)
print(APPROVAL_FILE)
print(MODEL_FILE)

