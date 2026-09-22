from pathlib import Path
from openpyxl import load_workbook
from collections import defaultdict, Counter
import csv
import json
import re
import sys


# ============================================================
# PATHS
# ============================================================

ROOT = Path(__file__).resolve().parents[2]

RAW_DIR = ROOT / "row_data"
DISCOVERY_DIR = ROOT / "outputs" / "discovery"
CONTRACT_DIR = ROOT / "outputs" / "contracts"
ARCH_DIR = ROOT / "docs" / "architecture"
VALIDATION_DIR = ROOT / "docs" / "validation"

CONTRACT_DIR.mkdir(parents=True, exist_ok=True)
ARCH_DIR.mkdir(parents=True, exist_ok=True)
VALIDATION_DIR.mkdir(parents=True, exist_ok=True)

FOUR_YEAR_FILE = DISCOVERY_DIR / "11_four_year_comparability_candidates.csv"
CORE_FILE = DISCOVERY_DIR / "12_core_portfolio_table_candidates.csv"
VALIDATION_FILE = DISCOVERY_DIR / "14_manual_validation_queue.csv"

OUTPUT_SELECTION = CONTRACT_DIR / "01_final_source_selection_candidates.csv"
OUTPUT_SIGNATURES = CONTRACT_DIR / "02_cross_year_structure_validation.csv"
OUTPUT_HEADERS = CONTRACT_DIR / "03_header_and_unit_evidence.csv"
OUTPUT_GRAIN = CONTRACT_DIR / "04_grain_contract_candidates.csv"
OUTPUT_REVIEW = CONTRACT_DIR / "05_manual_approval_register.csv"

DATA_CONTRACT_MD = ARCH_DIR / "DATA_CONTRACT_DRAFT.md"
VALIDATION_MD = VALIDATION_DIR / "PHASE_2A_VALIDATION_SUMMARY.md"

PREVIEW_ROWS = 25
MAX_COLS = 50


# ============================================================
# HELPERS
# ============================================================

def read_csv(path):
    if not path.exists():
        print(f"ERROR: Required file missing: {path}")
        sys.exit(1)

    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path, rows, fieldnames=None):

    if fieldnames is None:
        fieldnames = list(rows[0].keys()) if rows else []

    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=fieldnames
        )

        writer.writeheader()

        if rows:
            writer.writerows(rows)


def clean(value):

    if value is None:
        return ""

    text = str(value).strip()
    text = text.replace("\n", " ")
    text = text.replace("\r", " ")

    return re.sub(r"\s+", " ", text).strip()


def norm(value):

    text = clean(value).lower()

    replacements = {
        "–": "-",
        "—": "-",
        "_": " ",
        "/": " ",
        "\\": " ",
        "(": " ",
        ")": " ",
        ":": " ",
        ";": " ",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    text = re.sub(r"\s+", " ", text)

    return text.strip()


def safe_int(value):

    try:
        return int(float(value))
    except Exception:
        return None


def count_non_empty(row):
    return sum(
        1
        for value in row
        if clean(value)
    )


def numeric_like(value):

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return True

    text = clean(value)

    if not text:
        return False

    text = text.replace(",", "")
    text = text.replace("%", "")

    try:
        float(text)
        return True
    except Exception:
        return False


def text_like(value):

    return bool(
        clean(value)
        and not numeric_like(value)
    )


def row_signature(values):

    result = []

    for value in values:

        if not clean(value):
            result.append("B")

        elif numeric_like(value):
            result.append("N")

        else:
            result.append("T")

    return "".join(result)


def extract_source_year(filename):

    match = re.search(
        r"(2021|2022|2023|2024)",
        filename
    )

    if match:
        return int(match.group(1))

    return None


def classify_geography(text):

    t = norm(text)

    cluster_terms = [
        "health cluster",
        "cluster",
        "تجمع صحي",
        "التجمع الصحي",
        "تجمع"
    ]

    region_terms = [
        "health region",
        "health regions",
        "region",
        "regions",
        "المنطقة الصحية",
        "المناطق الصحية",
        "المنطقة",
        "المناطق"
    ]

    if any(norm(x) in t for x in cluster_terms):
        return "HEALTH_CLUSTER"

    if any(norm(x) in t for x in region_terms):
        return "HEALTH_REGION"

    return "UNKNOWN_OR_NATIONAL"


def detect_units(text):

    t = norm(text)

    units = []

    rules = {
        "COUNT": [
            "number",
            "no.",
            "count",
            "عدد"
        ],

        "PERCENT": [
            "%",
            "percent",
            "percentage",
            "نسبة",
            "بالمائة"
        ],

        "PER_1000": [
            "per 1,000",
            "per 1000",
            "لكل 1000",
            "لكل ألف"
        ],

        "PER_10000": [
            "per 10,000",
            "per 10000",
            "لكل 10000",
            "لكل عشرة آلاف"
        ],

        "RATE": [
            "rate",
            "معدل"
        ],

        "BEDS": [
            "bed",
            "beds",
            "سرير",
            "أسرة",
            "الأسرة"
        ]
    }

    for unit, terms in rules.items():

        if any(
            norm(term) in t
            for term in terms
        ):
            units.append(unit)

    return " | ".join(sorted(set(units)))


def detect_total_terms(text):

    t = norm(text)

    terms = [
        "total",
        "grand total",
        "subtotal",
        "الإجمالي",
        "المجموع",
        "جملة"
    ]

    hits = [
        term
        for term in terms
        if norm(term) in t
    ]

    return " | ".join(hits)


def detect_source_notes(text):

    t = norm(text)

    terms = [
        "source",
        "sources",
        "المصدر",
        "مصدر",
        "note",
        "notes",
        "ملاحظة",
        "ملاحظات"
    ]

    return " | ".join(
        term
        for term in terms
        if norm(term) in t
    )


def infer_grain_candidate(topic, geography, text):

    topic_text = norm(topic + " " + text)

    dimensions = ["Year"]

    if geography == "HEALTH_REGION":
        dimensions.append("Health Region")

    elif geography == "HEALTH_CLUSTER":
        dimensions.append("Health Cluster")

    else:
        dimensions.append("National / Other Geography")

    if "workforce" in topic_text:
        dimensions.append("Workforce Type / Profession")

    if "capacity" in topic_text:
        dimensions.append("Facility / Capacity Measure")

    if "activity" in topic_text:
        dimensions.append("Activity Type")

    if "population" in topic_text:
        dimensions.append("Population Measure")

    return " × ".join(dimensions)


# ============================================================
# LOAD DISCOVERY OUTPUTS
# ============================================================

four_year = read_csv(FOUR_YEAR_FILE)
core = read_csv(CORE_FILE)
manual_queue = read_csv(VALIDATION_FILE)


# Only core portfolio groups
core_groups = {
    row["comparison_group_id"]: row
    for row in core
}


candidate_rows = [
    row
    for row in four_year
    if row["comparison_group_id"] in core_groups
]


# ============================================================
# CACHE WORKBOOKS
# ============================================================

workbook_cache = {}


def get_workbook(filename):

    if filename in workbook_cache:
        return workbook_cache[filename]

    path = RAW_DIR / filename

    if not path.exists():
        raise FileNotFoundError(
            f"Workbook not found: {path}"
        )

    wb = load_workbook(
        path,
        read_only=True,
        data_only=False
    )

    workbook_cache[filename] = wb

    return wb


# ============================================================
# VALIDATE ACTUAL SOURCE SHEETS
# ============================================================

signature_rows = []
header_rows = []
grain_rows = []

group_evidence = defaultdict(list)


print("")
print("============================================================")
print("PHASE 2A — SOURCE CONTRACT VALIDATION")
print("============================================================")
print(f"Candidate source rows: {len(candidate_rows)}")
print("")


for index, row in enumerate(candidate_rows, start=1):

    group_id = row["comparison_group_id"]
    workbook_name = row["workbook_name"]
    sheet_name = row["sheet_name"]

    print(
        f"[{index}/{len(candidate_rows)}] "
        f"{group_id} | {row['source_year']} | {sheet_name}"
    )

    try:
        wb = get_workbook(workbook_name)

        if sheet_name not in wb.sheetnames:
            raise KeyError(
                f"Sheet not found: {sheet_name}"
            )

        ws = wb[sheet_name]

    except Exception as exc:

        signature_rows.append({
            "comparison_group_id": group_id,
            "source_year": row["source_year"],
            "workbook_name": workbook_name,
            "sheet_name": sheet_name,
            "status": "READ_ERROR",
            "error": str(exc)
        })

        continue


    max_row = ws.max_row or 0
    max_col = ws.max_column or 0

    scan_rows = min(PREVIEW_ROWS, max_row)
    scan_cols = min(MAX_COLS, max_col)

    preview = []

    for row_num, values in enumerate(
        ws.iter_rows(
            min_row=1,
            max_row=scan_rows,
            min_col=1,
            max_col=scan_cols,
            values_only=True
        ),
        start=1
    ):

        vals = list(values)

        preview.append(
            (row_num, vals)
        )


    # --------------------------------------------------------
    # Header candidate
    # --------------------------------------------------------

    probable_header = safe_int(
        row.get("probable_header_row")
    )

    probable_data_start = safe_int(
        row.get("probable_data_start_row")
    )

    if probable_header is None:
        probable_header = 1

    if probable_data_start is None:
        probable_data_start = probable_header + 1


    header_values = []

    if 1 <= probable_header <= len(preview):
        header_values = preview[
            probable_header - 1
        ][1]


    header_non_empty = [
        clean(v)
        for v in header_values
        if clean(v)
    ]


    # --------------------------------------------------------
    # Build contextual text
    # --------------------------------------------------------

    all_preview_values = []

    for _, values in preview:
        for value in values:

            value = clean(value)

            if value:
                all_preview_values.append(value)


    preview_text = " | ".join(
        all_preview_values
    )


    geography = classify_geography(
        preview_text
    )

    units = detect_units(
        preview_text
    )

    total_terms = detect_total_terms(
        preview_text
    )

    source_notes = detect_source_notes(
        preview_text
    )


    # --------------------------------------------------------
    # Structural signatures
    # --------------------------------------------------------

    first_data_signatures = []

    start_idx = max(
        probable_data_start - 1,
        0
    )

    for row_num, values in preview[start_idx:start_idx + 5]:

        first_data_signatures.append(
            f"R{row_num}:{row_signature(values)}"
        )


    # --------------------------------------------------------
    # Column evidence
    # --------------------------------------------------------

    header_rows.append({
        "comparison_group_id": group_id,
        "source_year": row["source_year"],
        "workbook_name": workbook_name,
        "sheet_name": sheet_name,
        "topic": row["topic"],
        "likely_table_title": row["likely_table_title"],
        "probable_header_row": probable_header,
        "probable_data_start_row": probable_data_start,
        "header_non_empty_count": len(header_non_empty),
        "header_values_json": json.dumps(
            header_non_empty,
            ensure_ascii=False
        ),
        "detected_geography": geography,
        "detected_units": units,
        "total_terms_found": total_terms,
        "source_or_note_terms": source_notes
    })


    signature_rows.append({
        "comparison_group_id": group_id,
        "source_year": row["source_year"],
        "workbook_name": workbook_name,
        "sheet_name": sheet_name,
        "status": "READ_OK",
        "max_rows": max_row,
        "max_columns": max_col,
        "header_row": probable_header,
        "data_start_row": probable_data_start,
        "header_field_count": len(header_non_empty),
        "first_data_row_signatures": " | ".join(
            first_data_signatures
        ),
        "detected_geography": geography,
        "detected_units": units,
        "error": ""
    })


    grain_candidate = infer_grain_candidate(
        row["topic"],
        geography,
        preview_text
    )


    grain_rows.append({
        "comparison_group_id": group_id,
        "source_year": row["source_year"],
        "topic": row["topic"],
        "workbook_name": workbook_name,
        "sheet_name": sheet_name,
        "grain_candidate": grain_candidate,
        "geography_candidate": geography,
        "units_candidate": units,
        "confirmed_grain": "",
        "confirmed_key_columns": "",
        "confirmed_measure_columns": "",
        "manual_status": "PENDING"
    })


    group_evidence[group_id].append({
        "year": safe_int(row["source_year"]),
        "rows": max_row,
        "cols": max_col,
        "header_count": len(header_non_empty),
        "geography": geography,
        "units": units,
        "topic": row["topic"],
        "title": row["likely_table_title"]
    })


# Close workbooks
for wb in workbook_cache.values():
    try:
        wb.close()
    except Exception:
        pass


# ============================================================
# CROSS-YEAR COMPARABILITY DECISION CANDIDATES
# ============================================================

selection_rows = []
approval_rows = []


for group_id, evidence in sorted(
    group_evidence.items()
):

    core_meta = core_groups[group_id]

    years = {
        e["year"]
        for e in evidence
        if e["year"] is not None
    }

    geographies = {
        e["geography"]
        for e in evidence
    }

    column_counts = {
        e["cols"]
        for e in evidence
    }

    header_counts = {
        e["header_count"]
        for e in evidence
    }

    units = {
        e["units"]
        for e in evidence
        if e["units"]
    }


    four_year_complete = years == {
        2021, 2022, 2023, 2024
    }


    same_geography = (
        len(geographies) == 1
        and "UNKNOWN_OR_NATIONAL" not in geographies
    )


    geography_review_required = (
        len(geographies) > 1
        or "UNKNOWN_OR_NATIONAL" in geographies
    )


    exact_column_structure = (
        len(column_counts) == 1
    )


    similar_header_width = (
        max(header_counts) - min(header_counts) <= 2
        if header_counts
        else False
    )


    unit_consistency_candidate = (
        len(units) <= 1
    )


    # Conservative recommendation only.
    # Never automatically approve.
    if (
        four_year_complete
        and same_geography
        and exact_column_structure
        and similar_header_width
        and unit_consistency_candidate
    ):
        contract_candidate = "STRONG_CANDIDATE"

    elif (
        four_year_complete
        and len(geographies) <= 2
    ):
        contract_candidate = "REVIEW_REQUIRED"

    else:
        contract_candidate = "HIGH_RISK_REVIEW"


    selection_rows.append({
        "comparison_group_id": group_id,
        "topic": core_meta["topic"],
        "representative_title":
            core_meta["representative_title"],
        "years_found":
            " | ".join(str(x) for x in sorted(years)),
        "four_year_complete":
            "YES" if four_year_complete else "NO",
        "geographies_found":
            " | ".join(sorted(geographies)),
        "same_geography_candidate":
            "YES" if same_geography else "NO",
        "column_counts":
            " | ".join(str(x) for x in sorted(column_counts)),
        "exact_column_structure":
            "YES" if exact_column_structure else "NO",
        "header_field_counts":
            " | ".join(str(x) for x in sorted(header_counts)),
        "similar_header_width":
            "YES" if similar_header_width else "NO",
        "units_found":
            " | ".join(sorted(units)),
        "unit_consistency_candidate":
            "YES" if unit_consistency_candidate else "NO",
        "contract_candidate_status":
            contract_candidate,

        # Human approval fields
        "same_business_definition_confirmed": "",
        "same_grain_confirmed": "",
        "same_geography_confirmed": "",
        "same_unit_definition_confirmed": "",
        "totals_handling_confirmed": "",
        "source_notes_reviewed": "",
        "final_source_decision": "",
        "decision_reason": ""
    })


    approval_rows.append({
        "comparison_group_id": group_id,
        "topic": core_meta["topic"],
        "representative_title":
            core_meta["representative_title"],
        "automated_candidate_status":
            contract_candidate,
        "business_definition":
            "PENDING",
        "grain":
            "PENDING",
        "geography":
            "PENDING",
        "units":
            "PENDING",
        "totals_subtotals":
            "PENDING",
        "footnotes_sources":
            "PENDING",
        "final_decision":
            "PENDING",
        "approved_fact_target": "",
        "review_notes": ""
    })


# ============================================================
# WRITE OUTPUTS
# ============================================================

write_csv(
    OUTPUT_SELECTION,
    selection_rows
)

write_csv(
    OUTPUT_SIGNATURES,
    signature_rows
)

write_csv(
    OUTPUT_HEADERS,
    header_rows
)

write_csv(
    OUTPUT_GRAIN,
    grain_rows
)

write_csv(
    OUTPUT_REVIEW,
    approval_rows
)


# ============================================================
# DATA CONTRACT DRAFT
# ============================================================

status_counts = Counter(
    row["contract_candidate_status"]
    for row in selection_rows
)

topic_counts = Counter(
    row["topic"]
    for row in selection_rows
)


with DATA_CONTRACT_MD.open(
    "w",
    encoding="utf-8"
) as f:

    f.write("# Data Contract Draft\n\n")

    f.write(
        "## Project\n\n"
        "**Saudi Healthcare Capacity & Performance Analytics 2021–2024**\n\n"
    )

    f.write(
        "## Status\n\n"
        "**DRAFT — NOT YET APPROVED**\n\n"
    )

    f.write(
        "This contract is generated from source discovery and structural "
        "validation. No source table is approved for transformation until "
        "business meaning, grain, geography, units, totals, and source notes "
        "are manually reviewed.\n\n"
    )

    f.write("## Source System\n\n")
    f.write(
        "- Official Saudi Ministry of Health Statistical Yearbooks\n"
        "- Years: 2021, 2022, 2023, 2024\n"
        "- Original Excel workbooks remain immutable\n\n"
    )

    f.write("## Candidate Analytical Domains\n\n")

    for topic, count in topic_counts.most_common():
        f.write(f"- {topic}: {count} candidate groups\n")

    f.write("\n## Structural Validation Status\n\n")

    for status, count in status_counts.items():
        f.write(f"- {status}: {count}\n")

    f.write(
        "\n## Geography Contract\n\n"
        "The project must distinguish at minimum:\n\n"
        "1. National-level observations\n"
        "2. Health Region observations\n"
        "3. Health Cluster observations\n\n"
        "`Health Region` and `Health Cluster` must not be merged into one "
        "entity without an explicit, validated mapping and compatible time "
        "logic.\n\n"
    )

    f.write(
        "## Grain Contract Rule\n\n"
        "Every final fact table must have a documented statement:\n\n"
        "`One row = <business dimensions>`\n\n"
        "No joins, aggregations, or KPI calculations may be approved before "
        "that grain is validated against the source tables.\n\n"
    )

    f.write(
        "## Preliminary Fact Candidates\n\n"
        "These are architecture candidates only:\n\n"
        "- Fact_PopulationIndicators\n"
        "- Fact_HealthcareCapacity\n"
        "- Fact_HealthcareWorkforce\n"
        "- Fact_HealthcareActivity\n\n"
    )

    f.write(
        "Final fact boundaries may change after manual source validation.\n\n"
    )

    f.write(
        "## Preliminary Shared Dimensions\n\n"
        "- Dim_Year\n"
        "- Dim_Geography\n"
        "- Dim_Indicator\n"
        "- Dim_FacilityType, if supported\n"
        "- Dim_WorkforceType / Profession, if supported\n"
        "- Dim_ActivityType, if supported\n\n"
    )

    f.write(
        "## Mandatory Validation Before Approval\n\n"
        "For every source group:\n\n"
        "- Same business definition across years\n"
        "- Same or explicitly reconcilable grain\n"
        "- Same geography concept\n"
        "- Same unit / denominator definition\n"
        "- Total and subtotal rows identified\n"
        "- Footnotes and methodological notes reviewed\n"
        "- Header and data boundaries confirmed\n"
        "- Final source inclusion decision documented\n\n"
    )

    f.write(
        "## Cleaning Status\n\n"
        "**NOT STARTED.**\n"
    )


# ============================================================
# VALIDATION SUMMARY
# ============================================================

with VALIDATION_MD.open(
    "w",
    encoding="utf-8"
) as f:

    f.write(
        "# Phase 2A — Data Contract Validation Summary\n\n"
    )

    f.write(
        f"- Core comparison groups inspected: "
        f"**{len(selection_rows)}**\n"
    )

    f.write(
        f"- Source-year sheets inspected: "
        f"**{len(signature_rows)}**\n"
    )

    for status, count in status_counts.most_common():
        f.write(
            f"- {status}: **{count}** groups\n"
        )

    f.write(
        "\n## Interpretation\n\n"
        "`STRONG_CANDIDATE` does not mean approved. "
        "It means only that automated structural checks found fewer "
        "differences. Human validation is still mandatory.\n\n"
    )

    f.write(
        "## Next Decision Gate\n\n"
        "Review `01_final_source_selection_candidates.csv` together with "
        "`03_header_and_unit_evidence.csv` and approve which groups enter "
        "the canonical dataset.\n\n"
    )

    f.write(
        "No cleaning, mapping, aggregation, or KPI calculation was "
        "performed in this phase.\n"
    )


# ============================================================
# FINAL OUTPUT
# ============================================================

print("")
print("============================================================")
print("PHASE 2A COMPLETE")
print("============================================================")
print(f"Core comparison groups inspected: {len(selection_rows)}")
print(f"Source-year sheets inspected:     {len(signature_rows)}")

for status, count in status_counts.most_common():
    print(f"{status:<28} {count}")

print("Source workbooks modified:        NO")
print("Cleaning performed:               NO")

print("")
print("UPLOAD THESE FILES:")
print(OUTPUT_SELECTION)
print(OUTPUT_SIGNATURES)
print(OUTPUT_HEADERS)
print(OUTPUT_GRAIN)
print(OUTPUT_REVIEW)
print(DATA_CONTRACT_MD)
print(VALIDATION_MD)

