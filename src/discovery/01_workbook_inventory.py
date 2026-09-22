from pathlib import Path
from openpyxl import load_workbook
import csv
import sys

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_DATA_DIR = PROJECT_ROOT / "row_data"
OUTPUT_DIR = PROJECT_ROOT / "outputs" / "tables"
OUTPUT_FILE = OUTPUT_DIR / "workbook_sheet_inventory.csv"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

def main():

    if not RAW_DATA_DIR.exists():
        print(f"ERROR: Raw data folder not found: {RAW_DATA_DIR}")
        sys.exit(1)

    excel_files = sorted([
        p for p in RAW_DATA_DIR.iterdir()
        if p.is_file()
        and p.suffix.lower() in {".xlsx", ".xlsm"}
        and not p.name.startswith("~$")
    ])

    if not excel_files:
        print(f"ERROR: No Excel files found in: {RAW_DATA_DIR}")
        sys.exit(1)

    records = []

    print("\n===== WORKBOOK DISCOVERY =====")

    for workbook_path in excel_files:

        print(f"\nWorkbook: {workbook_path.name}")

        try:
            workbook = load_workbook(
                workbook_path,
                read_only=False,
                data_only=False
            )

            for sheet_index, worksheet in enumerate(
                workbook.worksheets,
                start=1
            ):

                record = {
                    "workbook_name": workbook_path.name,
                    "sheet_index": sheet_index,
                    "sheet_name": worksheet.title,
                    "max_row": worksheet.max_row,
                    "max_column": worksheet.max_column,
                    "sheet_state": worksheet.sheet_state
                }

                records.append(record)

                print(
                    f"  {sheet_index:02d}. "
                    f"{worksheet.title} "
                    f"| Rows: {worksheet.max_row} "
                    f"| Columns: {worksheet.max_column} "
                    f"| State: {worksheet.sheet_state}"
                )

            workbook.close()

        except Exception as exc:
            print(f"  ERROR reading workbook: {exc}")

            records.append({
                "workbook_name": workbook_path.name,
                "sheet_index": "",
                "sheet_name": "",
                "max_row": "",
                "max_column": "",
                "sheet_state": f"ERROR: {exc}"
            })

    fieldnames = [
        "workbook_name",
        "sheet_index",
        "sheet_name",
        "max_row",
        "max_column",
        "sheet_state"
    ]

    with OUTPUT_FILE.open(
        mode="w",
        newline="",
        encoding="utf-8-sig"
    ) as file:

        writer = csv.DictWriter(
            file,
            fieldnames=fieldnames
        )

        writer.writeheader()
        writer.writerows(records)

    successful_sheets = [
        r for r in records
        if not str(r["sheet_state"]).startswith("ERROR")
    ]

    workbook_count = len(excel_files)
    sheet_count = len(successful_sheets)

    print("\n===== DISCOVERY SUMMARY =====")
    print(f"Workbooks discovered: {workbook_count}")
    print(f"Sheets discovered:    {sheet_count}")
    print(f"Output file:           {OUTPUT_FILE}")
    print("Source files modified: NO")
    print("==============================")


if __name__ == "__main__":
    main()
