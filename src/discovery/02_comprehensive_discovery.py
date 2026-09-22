from pathlib import Path
from openpyxl import load_workbook
from collections import Counter
import csv
import json
import re
import sys
import time


# ============================================================
# CONFIGURATION
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = PROJECT_ROOT / "row_data"

OUTPUT_DIR = PROJECT_ROOT / "outputs" / "discovery"
DOCS_DIR = PROJECT_ROOT / "docs" / "discovery"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
DOCS_DIR.mkdir(parents=True, exist_ok=True)

STRUCTURE_FILE = OUTPUT_DIR / "01_sheet_structural_profile.csv"
PREVIEW_FILE = OUTPUT_DIR / "02_sheet_previews.csv"
KEYWORDS_FILE = OUTPUT_DIR / "03_keyword_hits.csv"
REGION_FILE = OUTPUT_DIR / "04_region_candidates.csv"
YEAR_FILE = OUTPUT_DIR / "05_year_candidates.csv"
DQ_FILE = OUTPUT_DIR / "06_discovery_risk_register.csv"
SUMMARY_FILE = DOCS_DIR / "PHASE_1_DISCOVERY_SUMMARY.md"

PREVIEW_ROWS = 20
SCAN_ROWS = 40
SCAN_COLS = 40


# ============================================================
# DISCOVERY DICTIONARIES
# These are detection clues only — NOT cleaning mappings.
# ============================================================

YEAR_PATTERN = re.compile(r"\b(2021|2022|2023|2024)\b")

KEYWORD_GROUPS = {
    "region": [
        "region",
        "regions",
        "health region",
        "health regions",
        "المنطقة",
        "المناطق",
        "المنطقة الصحية",
        "المناطق الصحية"
    ],

    "year": [
        "year",
        "years",
        "السنة",
        "السنوات",
        "عام"
    ],

    "total": [
        "total",
        "grand total",
        "subtotal",
        "الإجمالي",
        "المجموع",
        "جملة"
    ],

    "source": [
        "source",
        "sources",
        "المصدر",
        "مصدر"
    ],

    "population": [
        "population",
        "populations",
        "السكان",
        "عدد السكان"
    ],

    "hospital": [
        "hospital",
        "hospitals",
        "مستشفى",
        "المستشفيات"
    ],

    "bed": [
        "bed",
        "beds",
        "سرير",
        "أسرة",
        "الأسرة"
    ],

    "physician": [
        "physician",
        "physicians",
        "doctor",
        "doctors",
        "طبيب",
        "أطباء",
        "الأطباء"
    ],

    "nurse": [
        "nurse",
        "nurses",
        "تمريض",
        "ممرض",
        "ممرضين"
    ],

    "activity": [
        "activity",
        "activities",
        "visits",
        "admissions",
        "operations",
        "الأنشطة",
        "النشاط",
        "الزيارات",
        "التنويم",
        "العمليات"
    ]
}


REGION_CLUES = [
    "riyadh",
    "makkah",
    "makkah al mukarramah",
    "makkah almukarramah",
    "madinah",
    "al madinah",
    "qassim",
    "al qassim",
    "eastern",
    "eastern region",
    "asir",
    "tabuk",
    "hail",
    "northern borders",
    "northern border",
    "jazan",
    "najran",
    "al baha",
    "baha",
    "al jouf",
    "jouf",

    "الرياض",
    "مكة",
    "مكة المكرمة",
    "المدينة",
    "المدينة المنورة",
    "القصيم",
    "الشرقية",
    "المنطقة الشرقية",
    "عسير",
    "تبوك",
    "حائل",
    "الحدود الشمالية",
    "جازان",
    "نجران",
    "الباحة",
    "الجوف"
]


# ============================================================
# HELPERS
# ============================================================

def safe_text(value):
    if value is None:
        return ""

    if isinstance(value, str):
        return value.strip()

    return str(value).strip()


def normalized_text(value):
    text = safe_text(value).lower()

    text = text.replace("\n", " ")
    text = text.replace("\r", " ")
    text = re.sub(r"\s+", " ", text)

    return text.strip()


def row_values(ws, row_number, max_cols):
    values = []

    for row in ws.iter_rows(
        min_row=row_number,
        max_row=row_number,
        min_col=1,
        max_col=max_cols,
        values_only=True
    ):
        values = [safe_text(v) for v in row]

    return values


def non_empty_count(values):
    return sum(1 for v in values if safe_text(v) != "")


def text_count(values):
    return sum(
        1
        for v in values
        if isinstance(v, str) and safe_text(v) != ""
    )


def numeric_count(values):
    count = 0

    for value in values:
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            count += 1

    return count


def detect_header_row(sample_rows):
    """
    Conservative heuristic only.
    Does NOT claim the row is definitively the real header.
    """

    candidates = []

    for row_number, values in sample_rows[:15]:
        populated = non_empty_count(values)

        if populated == 0:
            continue

        texts = text_count(values)
        numerics = numeric_count(values)

        score = (
            populated * 2
            + texts * 2
            - numerics
        )

        candidates.append(
            (score, row_number, populated, texts, numerics)
        )

    if not candidates:
        return None

    candidates.sort(reverse=True)

    return candidates[0][1]


def detect_first_non_empty_row(sample_rows):
    for row_number, values in sample_rows:
        if non_empty_count(values) > 0:
            return row_number

    return None


def detect_data_start(sample_rows, header_row):
    if header_row is None:
        return None

    for row_number, values in sample_rows:
        if row_number <= header_row:
            continue

        populated = non_empty_count(values)

        if populated < 2:
            continue

        numerics = numeric_count(values)
        texts = text_count(values)

        if numerics > 0 or texts >= 2:
            return row_number

    return None


def detect_keywords(workbook_name, sheet_name, row_number, col_number, value):
    hits = []

    text = normalized_text(value)

    if not text:
        return hits

    for group, words in KEYWORD_GROUPS.items():

        for word in words:
            if word.lower() in text:

                hits.append({
                    "workbook_name": workbook_name,
                    "sheet_name": sheet_name,
                    "row_number": row_number,
                    "column_number": col_number,
                    "keyword_group": group,
                    "matched_keyword": word,
                    "cell_value": safe_text(value)
                })

                break

    return hits


def detect_region_candidate(value):
    text = normalized_text(value)

    if not text:
        return None

    for region in REGION_CLUES:
        if region.lower() in text:
            return safe_text(value)

    return None


def detect_year_candidates(value):
    text = safe_text(value)

    if not text:
        return []

    return sorted(set(YEAR_PATTERN.findall(text)))


def classify_sheet_clues(keyword_counter):

    clues = []

    for key in [
        "region",
        "year",
        "population",
        "hospital",
        "bed",
        "physician",
        "nurse",
        "activity"
    ]:
        if keyword_counter.get(key, 0) > 0:
            clues.append(key)

    return ", ".join(clues)


# ============================================================
# MAIN DISCOVERY
# ============================================================

def main():

    start_time = time.time()

    if not RAW_DIR.exists():
        print(f"ERROR: Source folder not found: {RAW_DIR}")
        sys.exit(1)

    workbooks = sorted([
        p
        for p in RAW_DIR.iterdir()
        if p.is_file()
        and p.suffix.lower() in {".xlsx", ".xlsm"}
        and not p.name.startswith("~$")
    ])

    if not workbooks:
        print("ERROR: No Excel workbooks found.")
        sys.exit(1)


    structural_rows = []
    preview_rows = []
    keyword_rows = []
    region_rows = []
    year_rows = []
    dq_rows = []

    global_region_values = Counter()
    global_year_values = Counter()

    workbook_stats = []


    print("")
    print("============================================================")
    print("PHASE 1 — COMPREHENSIVE DATA DISCOVERY")
    print("============================================================")
    print(f"Source folder: {RAW_DIR}")
    print(f"Workbooks:     {len(workbooks)}")
    print("")


    # ========================================================
    # WORKBOOK LOOP
    # ========================================================

    for workbook_index, workbook_path in enumerate(workbooks, start=1):

        print(
            f"[{workbook_index}/{len(workbooks)}] "
            f"{workbook_path.name}"
        )

        workbook_sheet_count = 0

        try:
            wb = load_workbook(
                workbook_path,
                read_only=True,
                data_only=False
            )

        except Exception as exc:

            dq_rows.append({
                "severity": "CRITICAL",
                "workbook_name": workbook_path.name,
                "sheet_name": "",
                "risk_type": "WORKBOOK_READ_ERROR",
                "details": str(exc)
            })

            print(f"    ERROR: {exc}")
            continue


        # ====================================================
        # SHEET LOOP
        # ====================================================

        for ws in wb.worksheets:

            workbook_sheet_count += 1

            max_rows = ws.max_row or 0
            max_cols = ws.max_column or 0

            scan_max_rows = min(max_rows, SCAN_ROWS)
            scan_max_cols = min(max_cols, SCAN_COLS)

            sampled = []

            keyword_counter = Counter()

            sheet_regions = Counter()
            sheet_years = Counter()

            blank_rows_in_sample = 0
            sample_non_empty_cells = 0


            # ================================================
            # SAMPLE FIRST ROWS
            # ================================================

            for row_number, row in enumerate(
                ws.iter_rows(
                    min_row=1,
                    max_row=scan_max_rows,
                    min_col=1,
                    max_col=scan_max_cols,
                    values_only=True
                ),
                start=1
            ):

                values = list(row)

                populated = non_empty_count(values)

                if populated == 0:
                    blank_rows_in_sample += 1

                sample_non_empty_cells += populated

                sampled.append(
                    (row_number, values)
                )


                # Preview rows
                if row_number <= PREVIEW_ROWS:

                    preview_rows.append({
                        "workbook_name": workbook_path.name,
                        "sheet_name": ws.title,
                        "row_number": row_number,
                        "values_json": json.dumps(
                            [safe_text(v) for v in values],
                            ensure_ascii=False
                        )
                    })


                # Scan individual cells
                for col_number, value in enumerate(values, start=1):

                    if value is None:
                        continue

                    # Keyword detection
                    hits = detect_keywords(
                        workbook_path.name,
                        ws.title,
                        row_number,
                        col_number,
                        value
                    )

                    keyword_rows.extend(hits)

                    for hit in hits:
                        keyword_counter[
                            hit["keyword_group"]
                        ] += 1


                    # Year detection
                    detected_years = detect_year_candidates(value)

                    for year in detected_years:

                        sheet_years[year] += 1
                        global_year_values[year] += 1

                        year_rows.append({
                            "workbook_name": workbook_path.name,
                            "sheet_name": ws.title,
                            "row_number": row_number,
                            "column_number": col_number,
                            "year_candidate": year,
                            "cell_value": safe_text(value)
                        })


                    # Region candidate detection
                    region_candidate = detect_region_candidate(value)

                    if region_candidate:

                        sheet_regions[region_candidate] += 1
                        global_region_values[region_candidate] += 1

                        region_rows.append({
                            "workbook_name": workbook_path.name,
                            "sheet_name": ws.title,
                            "row_number": row_number,
                            "column_number": col_number,
                            "region_candidate": region_candidate
                        })


            # ================================================
            # STRUCTURAL INFERENCE
            # ================================================

            first_non_empty_row = detect_first_non_empty_row(sampled)
            probable_header_row = detect_header_row(sampled)
            probable_data_start = detect_data_start(
                sampled,
                probable_header_row
            )


            sheet_clues = classify_sheet_clues(
                keyword_counter
            )


            structural_rows.append({
                "workbook_name": workbook_path.name,
                "sheet_name": ws.title,
                "sheet_state": ws.sheet_state,
                "max_row_reported": max_rows,
                "max_column_reported": max_cols,
                "rows_scanned": scan_max_rows,
                "columns_scanned": scan_max_cols,
                "first_non_empty_row_candidate": first_non_empty_row or "",
                "probable_header_row": probable_header_row or "",
                "probable_data_start_row": probable_data_start or "",
                "blank_rows_in_sample": blank_rows_in_sample,
                "sample_non_empty_cells": sample_non_empty_cells,
                "region_keyword_hits": keyword_counter.get("region", 0),
                "year_keyword_hits": keyword_counter.get("year", 0),
                "total_keyword_hits": keyword_counter.get("total", 0),
                "source_keyword_hits": keyword_counter.get("source", 0),
                "region_candidate_count": sum(sheet_regions.values()),
                "year_candidate_count": sum(sheet_years.values()),
                "detected_years": ", ".join(sorted(sheet_years.keys())),
                "analytical_clues": sheet_clues
            })


            # ================================================
            # EARLY DQ / STRUCTURAL RISKS
            # ================================================

            if max_rows <= 2:

                dq_rows.append({
                    "severity": "INFO",
                    "workbook_name": workbook_path.name,
                    "sheet_name": ws.title,
                    "risk_type": "VERY_SMALL_SHEET",
                    "details": f"Reported rows = {max_rows}"
                })


            if max_cols <= 1:

                dq_rows.append({
                    "severity": "INFO",
                    "workbook_name": workbook_path.name,
                    "sheet_name": ws.title,
                    "risk_type": "SINGLE_COLUMN_OR_EMPTY_STRUCTURE",
                    "details": f"Reported columns = {max_cols}"
                })


            if first_non_empty_row and first_non_empty_row > 1:

                dq_rows.append({
                    "severity": "INFO",
                    "workbook_name": workbook_path.name,
                    "sheet_name": ws.title,
                    "risk_type": "LEADING_STRUCTURAL_ROWS",
                    "details": (
                        f"First non-empty row candidate = "
                        f"{first_non_empty_row}"
                    )
                })


            if blank_rows_in_sample >= 5:

                dq_rows.append({
                    "severity": "INFO",
                    "workbook_name": workbook_path.name,
                    "sheet_name": ws.title,
                    "risk_type": "MULTIPLE_BLANK_ROWS_IN_SAMPLE",
                    "details": (
                        f"{blank_rows_in_sample} blank rows "
                        f"within first {scan_max_rows} scanned rows"
                    )
                })


            if keyword_counter.get("total", 0) > 0:

                dq_rows.append({
                    "severity": "WARNING",
                    "workbook_name": workbook_path.name,
                    "sheet_name": ws.title,
                    "risk_type": "TOTAL_OR_SUBTOTAL_ROWS_POSSIBLE",
                    "details": (
                        f"Total-related keyword hits = "
                        f"{keyword_counter.get('total', 0)}"
                    )
                })


            if keyword_counter.get("source", 0) > 0:

                dq_rows.append({
                    "severity": "INFO",
                    "workbook_name": workbook_path.name,
                    "sheet_name": ws.title,
                    "risk_type": "SOURCE_OR_FOOTNOTE_TEXT_POSSIBLE",
                    "details": (
                        f"Source-related keyword hits = "
                        f"{keyword_counter.get('source', 0)}"
                    )
                })


        wb.close()


        workbook_stats.append({
            "workbook_name": workbook_path.name,
            "sheet_count": workbook_sheet_count
        })

        print(
            f"    Sheets profiled: {workbook_sheet_count}"
        )


    # ========================================================
    # WRITE CSV FILES
    # ========================================================

    def write_csv(path, rows, fieldnames):

        with path.open(
            "w",
            newline="",
            encoding="utf-8-sig"
        ) as f:

            writer = csv.DictWriter(
                f,
                fieldnames=fieldnames
            )

            writer.writeheader()
            writer.writerows(rows)


    write_csv(
        STRUCTURE_FILE,
        structural_rows,
        [
            "workbook_name",
            "sheet_name",
            "sheet_state",
            "max_row_reported",
            "max_column_reported",
            "rows_scanned",
            "columns_scanned",
            "first_non_empty_row_candidate",
            "probable_header_row",
            "probable_data_start_row",
            "blank_rows_in_sample",
            "sample_non_empty_cells",
            "region_keyword_hits",
            "year_keyword_hits",
            "total_keyword_hits",
            "source_keyword_hits",
            "region_candidate_count",
            "year_candidate_count",
            "detected_years",
            "analytical_clues"
        ]
    )


    write_csv(
        PREVIEW_FILE,
        preview_rows,
        [
            "workbook_name",
            "sheet_name",
            "row_number",
            "values_json"
        ]
    )


    write_csv(
        KEYWORDS_FILE,
        keyword_rows,
        [
            "workbook_name",
            "sheet_name",
            "row_number",
            "column_number",
            "keyword_group",
            "matched_keyword",
            "cell_value"
        ]
    )


    write_csv(
        REGION_FILE,
        region_rows,
        [
            "workbook_name",
            "sheet_name",
            "row_number",
            "column_number",
            "region_candidate"
        ]
    )


    write_csv(
        YEAR_FILE,
        year_rows,
        [
            "workbook_name",
            "sheet_name",
            "row_number",
            "column_number",
            "year_candidate",
            "cell_value"
        ]
    )


    write_csv(
        DQ_FILE,
        dq_rows,
        [
            "severity",
            "workbook_name",
            "sheet_name",
            "risk_type",
            "details"
        ]
    )


    # ========================================================
    # MARKDOWN SUMMARY
    # ========================================================

    total_sheets = len(structural_rows)

    sheets_with_region_clues = sum(
        1
        for r in structural_rows
        if r["region_candidate_count"] > 0
        or r["region_keyword_hits"] > 0
    )

    sheets_with_year_clues = sum(
        1
        for r in structural_rows
        if r["year_candidate_count"] > 0
        or r["year_keyword_hits"] > 0
    )

    sheets_with_total_clues = sum(
        1
        for r in structural_rows
        if r["total_keyword_hits"] > 0
    )

    risk_counts = Counter(
        r["risk_type"]
        for r in dq_rows
    )

    elapsed = time.time() - start_time


    with SUMMARY_FILE.open(
        "w",
        encoding="utf-8"
    ) as f:

        f.write(
            "# Phase 1 — Automated Data Discovery Summary\n\n"
        )

        f.write(
            "## Project\n\n"
            "Saudi Healthcare Capacity & Performance Analytics 2021–2024\n\n"
        )

        f.write(
            "## Discovery Scope\n\n"
            "This phase inspects the official source workbooks without "
            "cleaning, reshaping, joining, or modifying source data.\n\n"
        )

        f.write(
            "## Inventory\n\n"
        )

        f.write(
            f"- Workbooks discovered: **{len(workbooks)}**\n"
        )

        f.write(
            f"- Worksheets profiled: **{total_sheets}**\n"
        )

        f.write(
            f"- Sheets with region clues: **{sheets_with_region_clues}**\n"
        )

        f.write(
            f"- Sheets with year clues: **{sheets_with_year_clues}**\n"
        )

        f.write(
            f"- Sheets with total/subtotal clues: "
            f"**{sheets_with_total_clues}**\n\n"
        )


        f.write(
            "## Workbook Inventory\n\n"
        )

        for item in workbook_stats:
            f.write(
                f"- `{item['workbook_name']}`: "
                f"{item['sheet_count']} sheets\n"
            )


        f.write(
            "\n## Detected Year Candidates\n\n"
        )

        for value, count in global_year_values.most_common():
            f.write(
                f"- {value}: {count} sampled occurrences\n"
            )


        f.write(
            "\n## Most Frequent Region-like Source Values\n\n"
        )

        for value, count in global_region_values.most_common(40):
            f.write(
                f"- `{value}`: {count}\n"
            )


        f.write(
            "\n## Early Structural / Data Quality Risks\n\n"
        )

        if risk_counts:

            for risk_type, count in risk_counts.most_common():
                f.write(
                    f"- {risk_type}: {count}\n"
                )

        else:

            f.write(
                "- No automated discovery risks detected.\n"
            )


        f.write(
            "\n## Important Interpretation Rule\n\n"
            "Header rows, data-start rows, region values, and analytical "
            "clues in this phase are **candidates only**. "
            "They require manual validation before cleaning, joining, "
            "aggregation, or KPI implementation.\n"
        )


        f.write(
            "\n## Source Modification\n\n"
            "**NO source workbook was modified.**\n"
        )


    # ========================================================
    # FINAL SUMMARY
    # ========================================================

    print("")
    print("============================================================")
    print("PHASE 1 DISCOVERY COMPLETE")
    print("============================================================")

    print(f"Workbooks:                  {len(workbooks)}")
    print(f"Sheets profiled:            {total_sheets}")
    print(f"Sheets with region clues:   {sheets_with_region_clues}")
    print(f"Sheets with year clues:     {sheets_with_year_clues}")
    print(f"Risk records generated:     {len(dq_rows)}")
    print(f"Runtime seconds:            {elapsed:.1f}")
    print("Source files modified:      NO")

    print("")
    print("FILES CREATED:")
    print(f"1. {STRUCTURE_FILE}")
    print(f"2. {PREVIEW_FILE}")
    print(f"3. {KEYWORDS_FILE}")
    print(f"4. {REGION_FILE}")
    print(f"5. {YEAR_FILE}")
    print(f"6. {DQ_FILE}")
    print(f"7. {SUMMARY_FILE}")

    print("")
    print("PRIMARY FILE TO SHARE:")
    print(SUMMARY_FILE)

    print("")
    print("SUPPORTING FILE TO SHARE:")
    print(STRUCTURE_FILE)

    print("")
    print("PASS: Automated discovery completed successfully.")


if __name__ == "__main__":
    main()
